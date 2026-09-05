# 使い方: python3 budget/build_budget.py <第52期上期担当者.xlsx> [出力先.xlsx]
# 生成後は数式の再計算（Excelで開けば自動計算）を行ってください。
# -*- coding: utf-8 -*-
"""第52期上期担当者ファイル → 予算策定用ワークブック（全体＋部門別）を生成する。"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from collections import OrderedDict

import sys
SRC = sys.argv[1] if len(sys.argv) > 1 else '第52期上期担当者.xlsx'   # 元データ（20260905 1頁 シート）
OUT = sys.argv[2] if len(sys.argv) > 2 else 'budget/第52期上期担当者コード_予算策定.xlsx'

# ---------- 部門マッピング（第51期ファイルの部門別シート構成を踏襲） ----------
DEPT_CODES = OrderedDict([
    ('水産部X',        ['0010', '0017', '0031', '0101', '0102', '0104', '0106']),
    ('鳥栖・静岡・東京', ['0042', '0045', '0046', '0047', '0550', '0551', '0552']),
    ('松江',           ['0032', '0033', '0034', '0035', '0036', '0037', '0039']),
    ('境港',           ['0038', '0903', '0904', '0909', '0910', '0911', '0914', '0915', '0916',
                        '0918', '0919', '0930']),   # 0038 境その他 は名称どおり境港へ、0930 自治体(境港市) は新規コード
    ('下関',           ['0021', '0024', '0025', '0026', '0027', '0028', '0029']),  # 0028 自治体(下関市) は新規コード
    ('石見',           ['0011', '0012', '0013', '0014', '0015', '0016', '0018',
                        '0019', '0020', '0917']),
])
# 担当者名の読み替え（0918 は足立秀一氏退職のため「海外」として扱う）
RENAME = {'0918': '海外'}

# 全体シートで「部門」として表示する部門（0017 鮮魚は水産部X）
PRIMARY_DEPT = {}
for d, cs in DEPT_CODES.items():
    for c in cs:
        PRIMARY_DEPT.setdefault(c, d)

# ---------- 見出し ----------
PREV_MONTHS = ['2024/10', '2024/11', '2024/12', '2025/01', '2025/02', '2025/03']  # 前々期（第51期上期）
CUR_MONTHS  = ['2025/10', '2025/11', '2025/12', '2026/01', '2026/02', '2026/03']  # 前期（第52期上期）
HEADERS = (['得意先コード', '得意先略称', '担当者コード', '担当者名', '新担当者コード', '新担当者']
           + [m + '売上' for m in PREV_MONTHS] + ['前々期売上合計']
           + [m + '粗利益' for m in PREV_MONTHS] + ['前々期粗利益合計']
           + [m + '売上' for m in CUR_MONTHS] + ['前期売上合計']
           + [m + '粗利益' for m in CUR_MONTHS] + ['前期粗利益合計'])
assert len(HEADERS) == 34  # A..AH
NUM_COLS = list(range(7, 35))  # G..AH

# ---------- 書式 ----------
FONT = Font(name='游ゴシック', size=11)
FONT_B = Font(name='游ゴシック', size=11, bold=True)
FONT_GREEN = Font(name='游ゴシック', size=11, color='008000')
THIN = Side(style='thin')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill('solid', fgColor='CCFFFF')      # 元ファイルの見出し色（indexed 27）
INPUT_FILL = PatternFill('solid', fgColor='FFFF99')    # 入力欄
SUB_FILL = PatternFill('solid', fgColor='F2F2F2')      # 集計行
TOTAL_FILL = PatternFill('solid', fgColor='D9D9D9')    # 総計行
NUMFMT = '#,##0.0_ ;-#,##0.0_ ;"-"_ '
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

# ---------- 元データ読込 ----------
src = openpyxl.load_workbook(SRC).active
rows = []
for r in range(2, src.max_row + 1):
    a = src.cell(r, 1).value
    if not a:
        continue
    rec = {
        'code': str(a), 'name': src.cell(r, 2).value,
        'tcode': str(src.cell(r, 3).value), 'tname': src.cell(r, 4).value,
        'prev_sales': [src.cell(r, c).value or 0 for c in range(5, 11)],
        'prev_gp':    [src.cell(r, c).value or 0 for c in range(11, 17)],
        'cur_sales':  [src.cell(r, c).value or 0 for c in range(17, 23)],
        'cur_gp':     [src.cell(r, c).value or 0 for c in range(23, 29)],
    }
    rows.append(rec)
assert len({r['code'] for r in rows}) == len(rows), '得意先コードが重複しています'
unmapped = sorted({r['tcode'] for r in rows} - set(PRIMARY_DEPT))
assert not unmapped, f'部門未割当の担当者コード: {unmapped}'
rows.sort(key=lambda r: (r['tcode'], r['code']))

# 担当者マスタ（コード→名前、元データの出現順ではなくコード順）
master = OrderedDict()
for r in rows:
    master.setdefault(r['tcode'], r['tname'])
for _k, _v in RENAME.items():
    if _k in master:
        master[_k] = _v
master_active = master
MASTER_ROWS = 100  # マスタシートの参照範囲（新担当者追加用の余白込み）

wb = openpyxl.Workbook()
wb.remove(wb.active)

# 部門別シートの最終データ行を事前計算し、参照範囲を限定する
DEPT_LAST = {}
for _d, _cs in DEPT_CODES.items():
    DEPT_LAST[_d] = 1 + sum(1 for x in rows if x['tcode'] in _cs)


def style_header(ws, ncol, extra=None):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.font = FONT_B
        cell.fill = HDR_FILL
        cell.border = BORDER
        cell.alignment = CENTER
    ws.row_dimensions[1].height = 37.5


def set_widths(ws):
    widths = {'A': 15, 'B': 26.5, 'C': 13, 'D': 14.75, 'E': 15, 'F': 14.75}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    for c in NUM_COLS:
        ws.column_dimensions[L(c)].width = 12.5 if HEADERS[c - 1].endswith('合計') else 11


def write_data_row(ws, r, rec, new_code_formula=None):
    """1得意先分を書き込む。new_code_formula が None なら E は入力欄。"""
    ws.cell(r, 1, rec['code']).number_format = '@'
    ws.cell(r, 2, rec['name'])
    ws.cell(r, 3, rec['tcode']).number_format = '@'
    ws.cell(r, 4, RENAME.get(rec['tcode'], rec['tname']))
    e = ws.cell(r, 5)
    if new_code_formula is None:
        e.fill = INPUT_FILL
    else:
        e.value = new_code_formula
        e.font = FONT_GREEN
    e.number_format = '@'
    # F: 新担当者名はマスタから自動表示
    ws.cell(r, 6, f'=IF($E{r}="","",IFERROR(INDEX(担当者マスタ!$B$2:$B${MASTER_ROWS + 1},'
                  f'MATCH($E{r},担当者マスタ!$A$2:$A${MASTER_ROWS + 1},0)),"※マスタ未登録"))')
    vals = rec['prev_sales'] + [None] + rec['prev_gp'] + [None] + rec['cur_sales'] + [None] + rec['cur_gp'] + [None]
    for i, v in enumerate(vals):
        ws.cell(r, 7 + i, v)
    ws.cell(r, 13, f'=SUM(G{r}:L{r})')
    ws.cell(r, 20, f'=SUM(N{r}:S{r})')
    ws.cell(r, 27, f'=SUM(U{r}:Z{r})')
    ws.cell(r, 34, f'=SUM(AB{r}:AG{r})')
    for c in range(1, 35):
        cell = ws.cell(r, c)
        if cell.font != FONT_GREEN:
            cell.font = FONT
        cell.border = BORDER
        if c >= 7:
            cell.number_format = NUMFMT


# ================= 全体シート =================
ws_all = wb.create_sheet('全体')
for i, h in enumerate(HEADERS + ['部門'], 1):
    ws_all.cell(1, i, h)
style_header(ws_all, 35)
set_widths(ws_all)
ws_all.column_dimensions['AI'].width = 16
for i, rec in enumerate(rows):
    r = i + 2
    dept = PRIMARY_DEPT[rec['tcode']]
    # E: 部門別シートで入力した新担当者コードを自動参照
    dl = DEPT_LAST[dept]
    f = (f'=IFERROR(IF(INDEX(INDIRECT("\'"&$AI{r}&"\'!$E$2:$E${dl}"),MATCH($A{r},INDIRECT("\'"&$AI{r}&"\'!$A$2:$A${dl}"),0))="","",'
         f'INDEX(INDIRECT("\'"&$AI{r}&"\'!$E$2:$E${dl}"),MATCH($A{r},INDIRECT("\'"&$AI{r}&"\'!$A$2:$A${dl}"),0))),"")')
    write_data_row(ws_all, r, rec, new_code_formula=f)
    d = ws_all.cell(r, 35, dept)
    d.font = FONT
    d.border = BORDER
last_all = len(rows) + 1
tr = last_all + 1
ws_all.cell(tr, 3, '総計')
for c in NUM_COLS:
    ws_all.cell(tr, c, f'=SUBTOTAL(9,{L(c)}2:{L(c)}{last_all})')
for c in range(1, 36):
    cell = ws_all.cell(tr, c)
    cell.font = FONT_B
    cell.fill = TOTAL_FILL
    cell.border = BORDER
    if c >= 7:
        cell.number_format = NUMFMT
ws_all.freeze_panes = 'G2'
ws_all.auto_filter.ref = f'A1:AI{last_all}'
ws_all.print_area = f'A1:AH{tr}'
ws_all['E1'].comment = Comment('部門別シートのE列に入力した新担当者コードが自動で表示されます（このシートでは入力しません）。', 'Claude')
ws_all['AI1'].comment = Comment('担当者コードから判定した部門（0017 鮮魚は水産部X）。', 'Claude')

# ================= 部門別シート =================
dept_ranges = {}   # dept -> (first data row, last data row)
for dept, codes in DEPT_CODES.items():
    ws = wb.create_sheet(dept)
    for i, h in enumerate(HEADERS, 1):
        ws.cell(1, i, h)
    style_header(ws, 34)
    set_widths(ws)
    r = 2
    first_data = 2
    for code in codes:
        recs = [x for x in rows if x['tcode'] == code]
        if not recs:
            continue
        for rec in recs:
            write_data_row(ws, r, rec)
            r += 1
    last_data = r - 1
    ws.cell(r, 3, '総計')
    for c in NUM_COLS:
        ws.cell(r, c, f'=SUBTOTAL(9,{L(c)}{first_data}:{L(c)}{last_data})')
    for c in range(1, 35):
        cell = ws.cell(r, c)
        cell.font = FONT_B
        cell.fill = TOTAL_FILL
        cell.border = BORDER
        if c >= 7:
            cell.number_format = NUMFMT
    ws.freeze_panes = 'G2'
    ws.auto_filter.ref = f'A1:AH{last_data}'
    ws.print_area = f'A1:AH{r}'
    ws['E1'].comment = Comment('2026年10月以降の担当者コードをここに入力してください（黄色セル）。F列の氏名は担当者マスタから自動表示されます。', 'Claude')
    dept_ranges[dept] = (first_data, last_data)
    assert last_data == DEPT_LAST[dept], (dept, last_data, DEPT_LAST[dept])

# ================= 担当者マスタ =================
wm = wb.create_sheet('担当者マスタ')
for i, h in enumerate(['担当者コード', '担当者名', '部門', '備考'], 1):
    wm.cell(1, i, h)
style_header(wm, 4)
wm.column_dimensions['A'].width = 14
wm.column_dimensions['B'].width = 18
wm.column_dimensions['C'].width = 18
wm.column_dimensions['D'].width = 40
r = 2
for code, name in master_active.items():
    wm.cell(r, 1, code).number_format = '@'
    wm.cell(r, 2, name)
    wm.cell(r, 3, PRIMARY_DEPT[code])
    wm.cell(r, 4, '第52期上期データより')
    r += 1
first_blank = r
for rr in range(first_blank, MASTER_ROWS + 2):
    for c in range(1, 5):
        cell = wm.cell(rr, c)
        cell.fill = INPUT_FILL
    wm.cell(rr, 1).number_format = '@'
wm.cell(first_blank, 4, '← 新しい担当者はここから下（黄色）に コード・氏名・部門 を追加してください')
for rr in range(1, MASTER_ROWS + 2):
    for c in range(1, 5):
        cell = wm.cell(rr, c)
        cell.border = BORDER
        if rr > 1:
            cell.font = FONT
wm.freeze_panes = 'A2'

# ================= 担当者別集計（新担当者コード基準） =================
wsum = wb.create_sheet('担当者別集計')
SUM_HDR = ['新担当者コード', '新担当者名', '部門', '得意先数', '前々期売上', '前々期粗利', '前期売上', '前期粗利']
for i, h in enumerate(SUM_HDR, 1):
    wsum.cell(1, i, h)
style_header(wsum, len(SUM_HDR))
wsum.column_dimensions['A'].width = 15
wsum.column_dimensions['B'].width = 18
wsum.column_dimensions['C'].width = 18
for c in range(4, 9):
    wsum.column_dimensions[L(c)].width = 15
ER = f'全体!$E$2:$E${last_all}'
MR = f'全体!$M$2:$M${last_all}'
TR = f'全体!$T$2:$T${last_all}'
AAR = f'全体!$AA$2:$AA${last_all}'
AHR = f'全体!$AH$2:$AH${last_all}'
CNTFMT = '#,##0_ ;-#,##0_ ;"-"_ '
for i in range(MASTER_ROWS):
    r = i + 2
    m = i + 2  # 担当者マスタの行
    wsum.cell(r, 1, f'=IF(担当者マスタ!A{m}="","",担当者マスタ!A{m})').number_format = '@'
    wsum.cell(r, 2, f'=IF(担当者マスタ!A{m}="","",担当者マスタ!B{m})')
    wsum.cell(r, 3, f'=IF(担当者マスタ!A{m}="","",担当者マスタ!C{m})')
    wsum.cell(r, 4, f'=IF($A{r}="","",COUNTIF({ER},$A{r}))')
    wsum.cell(r, 5, f'=IF($A{r}="","",SUMIFS({MR},{ER},$A{r}))')
    wsum.cell(r, 6, f'=IF($A{r}="","",SUMIFS({TR},{ER},$A{r}))')
    wsum.cell(r, 7, f'=IF($A{r}="","",SUMIFS({AAR},{ER},$A{r}))')
    wsum.cell(r, 8, f'=IF($A{r}="","",SUMIFS({AHR},{ER},$A{r}))')
    for c in range(1, 9):
        cell = wsum.cell(r, c)
        cell.font = FONT
        cell.border = BORDER
        cell.number_format = NUMFMT if c >= 5 else (CNTFMT if c == 4 else cell.number_format)
last_sum = MASTER_ROWS + 1
r = last_sum + 1
wsum.cell(r, 1, '未設定')
wsum.cell(r, 2, '（新担当者コード空欄）')
wsum.cell(r, 4, f'=COUNTBLANK({ER})')
wsum.cell(r, 5, f'=SUMIFS({MR},{ER},"")')
wsum.cell(r, 6, f'=SUMIFS({TR},{ER},"")')
wsum.cell(r, 7, f'=SUMIFS({AAR},{ER},"")')
wsum.cell(r, 8, f'=SUMIFS({AHR},{ER},"")')
for c in range(1, 9):
    cell = wsum.cell(r, c)
    cell.font = FONT
    cell.fill = SUB_FILL
    cell.border = BORDER
    cell.number_format = NUMFMT if c >= 5 else CNTFMT
r += 1
wsum.cell(r, 1, '総計')
for c in range(4, 9):
    wsum.cell(r, c, f'=SUM({L(c)}2:{L(c)}{last_sum + 1})')
for c in range(1, 9):
    cell = wsum.cell(r, c)
    cell.font = FONT_B
    cell.fill = TOTAL_FILL
    cell.border = BORDER
    cell.number_format = NUMFMT if c >= 5 else CNTFMT
r += 1
wsum.cell(r, 1, '全体シート総計（検算）')
wsum.cell(r, 4, last_all - 1)
wsum.cell(r, 5, f'=全体!M{tr}')
wsum.cell(r, 6, f'=全体!T{tr}')
wsum.cell(r, 7, f'=全体!AA{tr}')
wsum.cell(r, 8, f'=全体!AH{tr}')
for c in range(1, 9):
    cell = wsum.cell(r, c)
    cell.font = FONT_GREEN
    cell.border = BORDER
    cell.number_format = NUMFMT if c >= 5 else CNTFMT
wsum.freeze_panes = 'D2'
check_row = r

# ================= 説明シート =================
wi = wb.create_sheet('説明', 0)
wi.column_dimensions['A'].width = 4
wi.column_dimensions['B'].width = 110
lines = [
    ('第52期上期 担当者コード（予算策定用）', True),
    ('', False),
    ('■ 元データ', True),
    ('・「20260905 1頁」シート（第52期上期担当者ファイル）の907得意先を、第51期上期担当者コードファイルと同じ構成で「全体」と部門別6シートに分割しています。', False),
    ('・今期（第53期）の予算策定用に、前々期 = 第51期上期（2024/10〜2025/03）、前期 = 第52期上期（2025/10〜2026/03）の実績を載せています。金額の単位は元データのまま（千円）です。', False),
    ('・各シートは 担当者コード → 得意先コード の順に並んでいます。合計・総計はすべて数式です。現担当者ごとの集計は設けていません（新担当者コードで集計します）。', False),
    ('', False),
    ('■ 2026年10月以降の担当者の設定方法', True),
    ('1. 各部門シート（水産部X／鳥栖・静岡・東京／松江／境港／下関／石見）の E列「新担当者コード」（黄色セル）に、新しい担当者コードを入力します。', False),
    ('2. F列「新担当者」は「担当者マスタ」シートから自動表示されます。「※マスタ未登録」と出た場合は担当者マスタにコードを追加してください。', False),
    ('3. 新しい担当者（新規コード）は「担当者マスタ」シートの黄色の空行に コード・氏名・部門 を追加します。', False),
    ('4. 「全体」シートの E列・F列は部門別シートの入力を自動で反映します（全体シートでは入力不要）。', False),
    ('5. 「担当者別集計」シートで、新担当者コードごとの得意先数・売上・粗利の合計を確認できます。「未設定」行が 0 になれば全得意先の割当完了です。', False),
    ('', False),
    ('■ 部門と担当者コードの対応（第51期ファイルの部門別シート構成を踏襲）', True),
]
for d, cs in DEPT_CODES.items():
    lines.append((f'・{d}：' + '、'.join(f'{c} {master[c]}' for c in cs if c in master), False))
lines += [
    ('・松江と境港は担当者コードの体系（003x＝松江、09xx＝境港）で分けています。0038 境その他 は名称どおり境港に置いています。', False),
    ('・0028 自治体（下関市）と 0930 自治体（境港市）は第52期で新たに登場したコードのため、それぞれ 下関・境港 に配置しています。', False),
    ('・0918 は足立秀一氏の退職に伴い、担当者名を「海外」に変更しています（担当者マスタも「0918 海外」）。', False),
    ('・0017 鮮魚 は水産部X のシートに掲載しています（石見には載せていません）。', False),
    ('', False),
    ('■ セルの色', True),
    ('・黄色 = 入力欄（部門別シートのE列、担当者マスタの空行）　・緑文字 = 他シート参照　・灰色 = 総計行', False),
]
for i, (t, b) in enumerate(lines, 1):
    c = wi.cell(i, 2, t)
    c.font = Font(name='游ゴシック', size=12 if i == 1 else 11, bold=b)
    c.alignment = Alignment(wrap_text=True, vertical='top')
wi.sheet_view.showGridLines = False

wb.active = 0
import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print('saved', OUT, 'rows', len(rows), 'dept ranges', dept_ranges, 'check_row', check_row)
