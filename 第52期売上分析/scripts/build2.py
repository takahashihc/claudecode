# -*- coding: utf-8 -*-
import json, re, copy
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree
from pptx.oxml.ns import qn

NS={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
V5=Path('in/v5x/ppt/charts')
def ser_vals(n):
    x=(V5/f'chart{n}.xml').read_text(encoding='utf-8'); out=[]
    for s in re.findall(r'<c:ser>.*?</c:ser>',x,re.S):
        vals=re.findall(r'<c:val>.*?</c:val>',s,re.S)[0]
        out.append([float(v) for v in re.findall(r'<c:v>(.*?)</c:v>',vals)])
    return out  # [単価, 数量]
aug=json.load(open('aug_result.json'))
GROUPS=['食品容器','フィルム・袋・ラミネート','紙箱・紙製品・ラベル・段ボール','消耗品','荷役','その他']
GCHART=[10,11,12,13,14,15]
OFFICES=['石見営業所','下関営業所','松江営業所','広域営業部','境港営業所','水産部']
OCHART=[16,17,18,19,20,21]
MONTHS=[4,5,6,7,8]
def series_for(names,charts):
    p,q=[],[]
    for nm,ch in zip(names,charts):
        u,v=ser_vals(ch)           # 4,5,6,7月
        a=aug[f'{nm}_8']
        p+= [round(x,4) for x in u]+[round(a['単価効果'],4)]
        q+= [round(x,4) for x in v]+[round(a['数量効果'],4)]
    return p,q
mon_u,mon_q=ser_vals(9); a8=aug['全社_8']
mon_u=[round(x,4) for x in mon_u]+[round(a8['単価効果'],4)]; mon_q=[round(x,4) for x in mon_q]+[round(a8['数量効果'],4)]
gp,gq=series_for(GROUPS,GCHART); op,oq=series_for(OFFICES,OCHART)

prs=Presentation('stage1.pptx')
S=prs.slides
def shapes_by_name(slide): return {sh.name:sh for sh in slide.shapes}
def set_text(shape, text):
    tf=shape.text_frame; p0=tf.paragraphs[0]
    for p in list(tf.paragraphs[1:]): p._p.getparent().remove(p._p)
    runs=p0.runs
    if not runs:
        p0.text=text; return
    runs[0].text=text
    for r in runs[1:]: r._r.getparent().remove(r._r)
def fix_labels(chart, thr, sz=None):
    cs=chart._chartSpace
    for ser in cs.iter(qn('c:ser')):
        dl=ser.find(qn('c:dLbls'))
        if dl is None: continue
        for d in dl.findall(qn('c:dLbl')): dl.remove(d)
        vals=[float(v.text) for v in ser.find(qn('c:val')).iter(qn('c:v'))]
        idxs=[i for i,v in enumerate(vals) if abs(v)<thr]
        for i in reversed(idxs):
            d=etree.SubElement(dl,qn('c:dLbl')); dl.remove(d); dl.insert(0,d)
            etree.SubElement(d,qn('c:idx')).set('val',str(i)); etree.SubElement(d,qn('c:delete')).set('val','1')
        if sz:
            for r in dl.iter(qn('a:defRPr')): r.set('sz',str(sz))
def replace_chart(chart, cats, u, q, multi=None, thr=0.02, sz=None):
    cd=CategoryChartData()
    if multi:
        for g in multi:
            c=cd.add_category(g)
            for m in MONTHS: c.add_sub_category(f'{m}月')
    else:
        cd.categories=cats
    cd.add_series('単価寄与',u,number_format='0.0%'); cd.add_series('数量寄与',q,number_format='0.0%')
    chart.replace_data(cd); fix_labels(chart,thr,sz)

# ---- slide6: 月次 1〜8月
s6=S[5]; sh=shapes_by_name(s6)
replace_chart(sh['Chart 8'].chart,[f'{m}月' for m in range(1,9)],mon_u,mon_q,thr=0.03)
set_text(sh['TextBox 5'],'機械類を除いた包装資材ベース。同一商品コードで突合し、単価寄与と数量寄与に分解（単価寄与＋数量寄与＝金額増減）。2025年→2026年の各月比較。8月を追加')
set_text(sh['TextBox 9'],'出所：202501～07　202601～07売上明細（1〜7月）、2508・2608タカハシ包装売上全明細（8月）。単価はラスパイレス価格指数（前年数量ウエイト）。単価が前年比0.25〜4.0倍の範囲外の品目は除外')
# ---- slide7: 商品群別 4〜8月
s7=S[6]; sh=shapes_by_name(s7)
replace_chart(sh['Chart 8'].chart,None,gp,gq,multi=GROUPS,thr=0.02,sz=700)
set_text(sh['TextBox 4'],'商品群別の単月推移 ―― 4〜8月の単価・数量寄与【機械類を除く】')
set_text(sh['TextBox 5'],'各月とも前年同月との比較。同一商品コードで突合し、単価寄与と数量寄与に分解（単価寄与＋数量寄与＝金額増減）。8月は単価主導が続き、食品容器・消耗品で数量減')
set_text(sh['TextBox 9'],'出所：202501～07　202601～07売上明細（4〜7月）、2508・2608タカハシ包装売上全明細（8月）を商品群_新旧対応表で分類。ラスパイレス価格指数（前年数量ウエイト）。単価が前年比0.25〜4.0倍の範囲外は除外。木箱は品目数不足のため非表示')
# ---- slide8: 営業所別 4〜8月
s8=S[7]; sh=shapes_by_name(s8)
replace_chart(sh['Chart 8'].chart,None,op,oq,multi=OFFICES,thr=0.02,sz=700)
set_text(sh['TextBox 4'],'営業所別の単月推移 ―― 4〜8月の単価・数量寄与【機械類を除く】')
set_text(sh['TextBox 5'],'各月とも前年同月との比較。営業所は担当者コードで突合。同一商品コードで単価寄与と数量寄与に分解。8月は全営業所で単価＋16〜24％、数量がプラスなのは水産部と広域のみ')
set_text(sh['TextBox 12'],'出所：202501～07　202601～07売上明細（4〜7月）、2508・2608タカハシ包装売上全明細（8月）。担当者コードで営業所に突合。ラスパイレス価格指数（前年数量ウエイト）。単価が前年比0.25〜4.0倍の範囲外は除外')

# ---- 新スライド: 8月単月の要因分解まとめ（slide6の見出し部品を複製）
layout=s6.slide_layout
ns=prs.slides.add_slide(layout)
for ph in list(ns.placeholders): ph._element.getparent().remove(ph._element)
src=shapes_by_name(s6)
for nm in ['Rectangle 1','Rectangle 2','TextBox 3','TextBox 4','TextBox 5','TextBox 7','TextBox 9']:
    el=copy.deepcopy(src[nm]._element); ns.shapes._spTree.append(el)
pic=src['Picture 6']
ns.shapes.add_picture('in/kx/ppt/media/image1.jpg',pic.left,pic.top,pic.width,pic.height)
nsh=shapes_by_name(ns)
set_text(nsh['TextBox 4'],'8月単月の要因分解 ―― 単価＋19.6％、数量▲2.5％【機械類を除く】')
set_text(nsh['TextBox 5'],'2025年8月と2026年8月を同一商品コードで突合し、単価寄与と数量寄与に分解。単位：千円（突合できた品目の売上合計）')
set_text(nsh['TextBox 9'],'出所：2508・2608タカハシ包装売上全明細。ラスパイレス価格指数（前年数量ウエイト）。単価が前年比0.25〜4.0倍の範囲外の品目は除外。木箱は品目数不足のため非表示。営業所は担当者コードで突合')

NAVY=RGBColor(0x1F,0x38,0x64); CHAR=RGBColor(0x40,0x40,0x40); GREY=RGBColor(0xF2,0xF2,0xF2); WHITE=RGBColor(0xFF,0xFF,0xFF)
GOLD=RGBColor(0xB8,0x86,0x0B); BLUE=RGBColor(0x44,0x72,0xC4); RED=RGBColor(0xC0,0x00,0x00)
def cell_text(cell,text,size=9,bold=False,color=CHAR,align=PP_ALIGN.RIGHT,fill=None):
    cell.text=''; tf=cell.text_frame; p=tf.paragraphs[0]; r=p.add_run(); r.text=text
    r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name='游ゴシック'; p.alignment=align
    cell.margin_left=cell.margin_right=Inches(0.05); cell.margin_top=cell.margin_bottom=Inches(0.02)
    cell.vertical_anchor=MSO_ANCHOR.MIDDLE
    if fill is not None:
        cell.fill.solid(); cell.fill.fore_color.rgb=fill
    else:
        cell.fill.background()
def pct(v): return ('＋' if v>=0 else '▲')+f'{abs(v)*100:.1f}％'
def add_table(slide,x,y,w,title,names,key_total):
    rows=[(nm,aug.get(f'{nm}_8')) for nm in names]+[('全社（機械類除く）',aug['全社_8'])]
    rows=[r for r in rows if r[1]]
    hdr=['区分','突合\n品目数','前年8月\n売上','当年8月\n売上','金額増減','単価寄与','数量寄与']
    tb=slide.shapes.add_textbox(x,y,w,Inches(0.3)); tf=tb.text_frame; tf.margin_left=0; p=tf.paragraphs[0]; r=p.add_run(); r.text=title
    r.font.size=Pt(12); r.font.bold=True; r.font.color.rgb=NAVY; r.font.name='游ゴシック'
    n=len(rows)+1
    gs=slide.shapes.add_table(n,7,x,y+Inches(0.32),w,Inches(0.3)*n)
    t=gs.table
    tblPr=t._tbl.tblPr; tblPr.set('bandRow','0'); tblPr.set('firstRow','0')
    st=tblPr.find(qn('a:tableStyleId'))
    if st is not None: tblPr.remove(st)
    cw=[Inches(1.9),Inches(0.6),Inches(0.85),Inches(0.85),Inches(0.62),Inches(0.62),Inches(0.62)]
    tot=sum(cw); cw=[Emu(int(c*w/tot)) for c in cw]
    for j,c in enumerate(cw): t.columns[j].width=c
    for j,h in enumerate(hdr): cell_text(t.cell(0,j),h,8,True,WHITE,PP_ALIGN.CENTER,NAVY)
    t.rows[0].height=Inches(0.42)
    for i,(nm,v) in enumerate(rows,1):
        last=(i==len(rows)); fill=RGBColor(0xE8,0xEE,0xF7) if last else (GREY if i%2==0 else None)
        cell_text(t.cell(i,0),nm,9,last,CHAR,PP_ALIGN.LEFT,fill)
        cell_text(t.cell(i,1),f"{v['品目数']:,}",9,last,CHAR,PP_ALIGN.RIGHT,fill)
        cell_text(t.cell(i,2),f"{v['前年売上']/1000:,.0f}",9,last,CHAR,PP_ALIGN.RIGHT,fill)
        cell_text(t.cell(i,3),f"{v['当年売上']/1000:,.0f}",9,last,CHAR,PP_ALIGN.RIGHT,fill)
        cell_text(t.cell(i,4),pct(v['金額増減']),9,last,CHAR,PP_ALIGN.RIGHT,fill)
        cell_text(t.cell(i,5),pct(v['単価効果']),9,True,GOLD if v['単価効果']>=0 else RED,PP_ALIGN.RIGHT,fill)
        cell_text(t.cell(i,6),pct(v['数量効果']),9,True,BLUE if v['数量効果']>=0 else RED,PP_ALIGN.RIGHT,fill)
        t.rows[i].height=Inches(0.3)
    return gs
add_table(ns,Inches(0.45),Inches(1.25),Inches(6.1),'商品群別（機械類を除く）',GROUPS,None)
add_table(ns,Inches(6.78),Inches(1.25),Inches(6.1),'営業所別（機械類を除く）',OFFICES,None)

# 3つのポイント
def card(slide,x,y,w,h,head,body,accent):
    r=slide.shapes.add_shape(1,x,y,w,h); r.fill.solid(); r.fill.fore_color.rgb=RGBColor(0xF7,0xF9,0xFC); r.line.fill.background(); r.shadow.inherit=False
    t=slide.shapes.add_textbox(x+Inches(0.15),y+Inches(0.1),w-Inches(0.3),h-Inches(0.2)); tf=t.text_frame; tf.word_wrap=True; tf.margin_left=tf.margin_right=0
    p=tf.paragraphs[0]; rr=p.add_run(); rr.text=head; rr.font.size=Pt(12); rr.font.bold=True; rr.font.color.rgb=accent; rr.font.name='游ゴシック'
    p2=tf.add_paragraph(); p2.space_before=Pt(4); rr=p2.add_run(); rr.text=body; rr.font.size=Pt(10); rr.font.color.rgb=CHAR; rr.font.name='游ゴシック'
def pc(k,f): return pct(aug[k][f])
y0=Inches(4.15); h=Inches(2.05); w=Inches(4.0); gap=Inches(0.2)
card(ns,Inches(0.45),y0,w,h,'① 8月も単価主導。数量は2か月連続でマイナス',
     f"突合品目ベースの金額増減{pc('全社_8','金額増減')}のうち、単価寄与{pc('全社_8','単価効果')}・数量寄与{pc('全社_8','数量効果')}。7月（単価＋19.0％・数量▲1.2％）に続き、単価上昇が売上を押し上げる一方で数量は弱い。",GOLD)
card(ns,Inches(0.45)+w+gap,y0,w,h,'② 食品容器は単価＋23.9％だが数量▲7.6％',
     f"フィルム・袋は単価{pc('フィルム・袋・ラミネート_8','単価効果')}に数量{pc('フィルム・袋・ラミネート_8','数量効果')}が上乗せ。紙箱・段ボールは単価{pc('紙箱・紙製品・ラベル・段ボール_8','単価効果')}と転嫁が緩やか。消耗品は数量{pc('消耗品_8','数量効果')}、荷役は単価{pc('荷役_8','単価効果')}と唯一の単価マイナス。",NAVY)
card(ns,Inches(0.45)+2*(w+gap),y0,w,h,'③ 数量がプラスの営業所は水産部と広域のみ',
     f"水産部は単価{pc('水産部_8','単価効果')}・数量{pc('水産部_8','数量効果')}と数量を伴う伸び。松江は単価{pc('松江営業所_8','単価効果')}で最大。石見{pc('石見営業所_8','数量効果')}・境港{pc('境港営業所_8','数量効果')}・下関{pc('下関営業所_8','数量効果')}と数量減が目立つ。",BLUE)
note=ns.shapes.add_textbox(Inches(0.45),Inches(6.35),Inches(12.4),Inches(0.6)); tf=note.text_frame; tf.word_wrap=True; tf.margin_left=0
p=tf.paragraphs[0]; r=p.add_run()
r.text='※ 包装資材計（機械類除く）の8月売上は前年比＋7.5％（240.4→258.5百万円）。同一品目で突合できた分は＋17.0％で、差は品目の入替（新規・終売）による。全社売上は機械類（前年8月87.5百万円→8.1百万円）の反動で▲18.7％。'
r.font.size=Pt(9); r.font.color.rgb=CHAR; r.font.name='游ゴシック'

# 新スライドを slide8（営業所別）の直後へ移動
lst=prs.slides._sldIdLst; ids=list(lst); new=ids[-1]; lst.remove(new); lst.insert(8,new)
# ページ番号を通し番号に
for i,s in enumerate(prs.slides,1):
    for shp in s.shapes:
        if shp.has_text_frame and '社外秘　P' in shp.text_frame.text:
            set_text(shp,f'㈱タカハシ包装センター　社外秘　P{i}')
prs.save('stage2.pptx'); print('stage2 saved', len(prs.slides))
