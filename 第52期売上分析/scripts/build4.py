# -*- coding: utf-8 -*-
import copy, pandas as pd, numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
NAVY=RGBColor(0x1F,0x38,0x64); CHAR=RGBColor(0x40,0x40,0x40); GREY=RGBColor(0xF2,0xF2,0xF2); WHITE=RGBColor(0xFF,0xFF,0xFF)
GOLD=RGBColor(0xB8,0x86,0x0B); BLUE=RGBColor(0x44,0x72,0xC4); RED=RGBColor(0xC0,0x00,0x00); LIGHT=RGBColor(0xF7,0xF9,0xFC); PALE=RGBColor(0xE8,0xEE,0xF7); GREEN=RGBColor(0x2E,0x7D,0x32)
prs=Presentation('stage3.pptx'); S=prs.slides
def by_name(sl): return {sh.name:sh for sh in sl.shapes}
def set_text(shape,text):
    tf=shape.text_frame; p0=tf.paragraphs[0]
    for p in list(tf.paragraphs[1:]): p._p.getparent().remove(p._p)
    runs=p0.runs
    if not runs: p0.text=text; return
    runs[0].text=text
    for r in runs[1:]: r._r.getparent().remove(r._r)
def run(p,text,size=10,bold=False,color=CHAR):
    r=p.add_run(); r.text=text; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name='游ゴシック'; return r
def card(slide,x,y,w,h,head,paras,accent):
    r=slide.shapes.add_shape(1,x,y,w,h); r.fill.solid(); r.fill.fore_color.rgb=LIGHT; r.line.fill.background(); r.shadow.inherit=False
    t=slide.shapes.add_textbox(x+Inches(0.15),y+Inches(0.08),w-Inches(0.3),h-Inches(0.16)); tf=t.text_frame; tf.word_wrap=True; tf.margin_left=tf.margin_right=0
    run(tf.paragraphs[0],head,12,True,accent)
    for para in paras:
        p=tf.add_paragraph(); p.space_before=Pt(4)
        if isinstance(para,tuple): run(p,para[0],**para[1])
        else: run(p,para,10)
def cell(c,text,size=9,bold=False,color=CHAR,align=PP_ALIGN.RIGHT,fill=None):
    c.text=''; p=c.text_frame.paragraphs[0]; run(p,text,size,bold,color); p.alignment=align
    c.margin_left=c.margin_right=Inches(0.05); c.margin_top=c.margin_bottom=Inches(0.02); c.vertical_anchor=MSO_ANCHOR.MIDDLE
    if fill is not None: c.fill.solid(); c.fill.fore_color.rgb=fill
    else: c.fill.background()
def table(slide,x,y,w,widths,rows,hdr,row_h=0.28,hdr_h=0.36,size=9,bold_last=False):
    n=len(rows)+1; gs=slide.shapes.add_table(n,len(hdr),x,y,w,Inches(row_h)*n); t=gs.table
    tp=t._tbl.tblPr; tp.set('bandRow','0'); tp.set('firstRow','0'); st=tp.find(qn('a:tableStyleId'))
    if st is not None: tp.remove(st)
    tot=sum(widths)
    for j,cw in enumerate(widths): t.columns[j].width=Emu(int(w*cw/tot))
    for j,h in enumerate(hdr): cell(t.cell(0,j),h,8,True,WHITE,PP_ALIGN.CENTER,NAVY)
    t.rows[0].height=Inches(hdr_h)
    for i,row in enumerate(rows,1):
        last=bold_last and i==len(rows); fill=PALE if last else (GREY if i%2==0 else None)
        for j,v in enumerate(row):
            txt,kw=(v if isinstance(v,tuple) else (v,{}))
            cell(t.cell(i,j),txt,kw.get('size',size),kw.get('bold',last),kw.get('color',CHAR),kw.get('align',PP_ALIGN.LEFT if j==0 else PP_ALIGN.RIGHT),fill)
        t.rows[i].height=Inches(row_h)
    return gs
def pct(v,d=1): return ('＋' if v>=0 else '▲')+f'{abs(v)*100:.{d}f}％'
def ptv(v,d=1): return ('＋' if v>=0 else '▲')+f'{abs(v)*100:.{d}f}pt'
def title_box(slide,x,y,w,text):
    tb=slide.shapes.add_textbox(x,y,w,Inches(0.3)); tb.text_frame.margin_left=0; run(tb.text_frame.paragraphs[0],text,11,True,NAVY)

# テンプレ：P10（ブリッジ）から見出し部品を複製
src=S[9]; ns=prs.slides.add_slide(src.slide_layout)
for ph in list(ns.placeholders): ph._element.getparent().remove(ph._element)
sh=by_name(src)
for nm in ['Rectangle 1','Rectangle 2','TextBox 3','TextBox 4','TextBox 5','TextBox 7','TextBox 9']: ns.shapes._spTree.append(copy.deepcopy(sh[nm]._element))
pic=[s for s in src.shapes if s.shape_type==13][0]; ns.shapes.add_picture('in/kx/ppt/media/image1.jpg',pic.left,pic.top,pic.width,pic.height)
n=by_name(ns)
set_text(n['TextBox 4'],'分類別の平均単価で見たスペック変更の影響【機械類を除く】')
set_text(n['TextBox 5'],'分類ごとの平均単価（売上金額÷売上数量）は突合できない品目も含むため、仕様変更で商品コードが変わった分も反映される。前年売上ウエイトで加重。8月の前年同月比')
set_text(n['TextBox 9'],'出所：2508・2608タカハシ包装売上全明細。分類は旧43分類（分類コード1）。「散らばり」は分類内の販売単価の上位10％÷下位10％。散らばり10倍以内を「単位がそろう」と判定')

g=pd.read_csv('uv_by_group.csv'); c=pd.read_csv('uv_by_class.csv')
# 左：商品群別
title_box(ns,Inches(0.45),Inches(1.2),Inches(6),'商品群別 ―― 分類別平均単価指数 と 突合単価効果 の比較')
rows=[]
for _,r in g.iterrows():
    if r['商品群']=='木箱': continue
    nm='全体（機械類除く）' if r['商品群']=='全体' else r['商品群']
    d=r['差']
    rows.append([nm,pct(r['売上増減']),(pct(r['分類別平均単価指数']),{'bold':True,'color':GOLD}),pct(r['突合単価効果']),(ptv(d),{'bold':True,'color':RED if d<0 else GREEN})])
table(ns,Inches(0.45),Inches(1.52),Inches(6.1),[2.2,1,1.2,1.1,1.1],rows,['商品群','売上増減','分類別\n平均単価指数','突合\n単価効果','差（スペック・\nミックス）'],row_h=0.29,hdr_h=0.42,bold_last=True)
# 右：分類別（単位がそろう分類 / そろわない分類）
title_box(ns,Inches(6.78),Inches(1.2),Inches(6.1),'主要分類 ―― 平均単価が使える分類と、使えない分類')
homog=[2,1,3,32,27,26]; hetero=[5,4,6,99]
rows=[]
for code in homog+hetero:
    r=c[c['分類']==code].iloc[0]; disp=10**r['単価ばらつき']; ok=code in homog
    d=r['平均単価変化']-r['突合単価効果'] if not np.isnan(r['突合単価効果']) else np.nan
    rows.append([r['名称'],pct(r['平均単価変化']),pct(r['突合単価効果']) if not np.isnan(r['突合単価効果']) else '－',
                 (ptv(d) if not np.isnan(d) else '－',{'bold':True,'color':RED if (not np.isnan(d) and d<0) else GREEN}),
                 f'{disp:,.0f}倍',(('○ 平均単価' if ok else '× 突合のみ'),{'bold':True,'color':GREEN if ok else RED,'align':PP_ALIGN.CENTER})])
table(ns,Inches(6.78),Inches(1.52),Inches(6.1),[1.7,1.05,1.0,1.05,0.9,1.1],rows,['分類','平均単価\n変化','突合\n単価効果','差','単価の\n散らばり','判定'],row_h=0.25,hdr_h=0.38)
# 下：3カード
y=Inches(4.85); h=Inches(1.7); w=Inches(4.0); gap=Inches(0.2)
ga=g[g['商品群']=='全体'].iloc[0]; tr=c[c['分類']==2].iloc[0]; fo=c[c['分類']==1].iloc[0]
card(ns,Inches(0.45),y,w,h,f"① スペック変更の影響は全体で{ptv(ga['差'])}",
     [f"突合{pct(ga['突合単価効果'])}に対し、分類別平均単価では{pct(ga['分類別平均単価指数'])}。仕様変更や安価品への切替が値上げの一部を相殺している。",('→ 計画の単価前提は突合の数字より低めに置く。',{'size':10,'bold':True,'color':NAVY})],GOLD)
card(ns,Inches(0.45)+w+gap,y,w,h,'② トレーはスペックダウンが大きい',
     [f"トレーは突合{pct(tr['突合単価効果'])}に対し平均単価{pct(tr['平均単価変化'])}で、差{ptv(tr['平均単価変化']-tr['突合単価効果'])}。安価なトレーへの切替・サイズダウンが進んでいる。発泡スチロールは逆に高単価品へシフト（{ptv(fo['平均単価変化']-fo['突合単価効果'])}）。"],RED)
card(ns,Inches(0.45)+2*(w+gap),y,w,h,'③ 使い分け（ハイブリッド）',
     ['単位がそろう分類（トレー・発泡・段ボール等）は平均単価を採用し、スペック変更込みの単価変化とする。',
      '枚売りとケース売りが混在する分類（袋類・フィルム・シール・その他）は突合で測り、スペック変更分は品目入替（P10）で別建て。'],BLUE)
# 挿入位置：P10の直後
lst=prs.slides._sldIdLst; ids=list(lst); new=ids[-1]; lst.remove(new); lst.insert(10,new)
for i,sl in enumerate(prs.slides,1):
    for shp in sl.shapes:
        if shp.has_text_frame and '社外秘　P' in shp.text_frame.text: set_text(shp,f'㈱タカハシ包装センター　社外秘　P{i}')
# P10 の右下カードの参照を更新
for shp in S[9].shapes:
    if shp.has_text_frame and '計画に使うのは' in shp.text_frame.text:
        for p in shp.text_frame.paragraphs[1:]:
            for r in p.runs: r.text=r.text.replace('単価寄与は「値上げがどこまで進んだか」を測る指標として使う。','スペック変更の影響は次ページで分類別平均単価から確認する。')
prs.save('stage4.pptx'); print('stage4',len(prs.slides))
