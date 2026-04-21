import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.styles.numbers import FORMAT_DATE_DATETIME

# ─── カラー定義 ───────────────────────────────────────────────
DARK_BLUE   = "1F4E79"
MID_BLUE    = "2E75B6"
LIGHT_BLUE  = "BDD7EE"
PALE_BLUE   = "DEEAF1"
WHITE       = "FFFFFF"
ORANGE_HDR  = "C55A11"   # 中東・ホルムズ列ヘッダ強調
PALE_ORANGE = "FCE4D6"   # 中東・ホルムズ列データ背景
GREEN_HDR   = "375623"
PALE_GREEN  = "E2EFDA"
GRAY_FILL   = "F2F2F2"
RED_URGENT  = "FF0000"
YELLOW_MID  = "FFCC00"

def fill(color):
    return PatternFill(start_color=color, end_color=color, fill_type="solid")

def font(color=WHITE, bold=False, size=10, name="Meiryo"):
    return Font(name=name, bold=bold, color=color, size=size)

def border(style="thin"):
    s = Side(style=style)
    return Border(left=s, right=s, top=s, bottom=s)

def align(h="center", v="center", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

# ─── ワークブック作成 ──────────────────────────────────────────
wb = openpyxl.Workbook()

# ════════════════════════════════════════════════════════════════
# シート①：メインニュース収集フォーマット
# ════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = "ニュース収集フォーマット"

# 商品群（行）
products = [
    ("段ボール",               "古紙・クラフト紙"),
    ("発泡スチロール",          "スチレンモノマー(石化)"),
    ("食品トレー",              "PS・PP樹脂(石化)"),
    ("ボードン袋",              "PE樹脂(石化)"),
    ("ストレッチフィルム",      "LLDPE樹脂(石化)"),
    ("ニトリル手袋",            "アクリロニトリル・ブタジエン(石化)"),
    ("プラスチックコンテナー",  "PP・HDPE樹脂(石化)"),
]

# ニュース収集列定義
# (ヘッダ表示名, ヘッダ塗り色, データ塗り色_偶数行, データ塗り色_奇数行, 列幅)
columns = [
    # 基本情報
    ("商品群",            DARK_BLUE,  MID_BLUE,    MID_BLUE,    18),
    ("主要原材料",        DARK_BLUE,  PALE_BLUE,   WHITE,       20),
    # 原材料・調達
    ("原材料\n価格動向",  "37618A",   PALE_BLUE,   WHITE,       18),
    ("国内仕入\n価格動向","37618A",   PALE_BLUE,   WHITE,       18),
    ("主要\n調達先・産地","37618A",   PALE_BLUE,   WHITE,       20),
    ("代替調達先\n・素材", "37618A",  PALE_BLUE,   WHITE,       20),
    # 中東・ホルムズ ★強調★
    ("ホルムズ海峡\n通行状況", ORANGE_HDR, PALE_ORANGE, PALE_ORANGE, 20),
    ("中東情勢\n直接影響",     ORANGE_HDR, PALE_ORANGE, PALE_ORANGE, 20),
    ("原油・ナフサ\n価格",     ORANGE_HDR, PALE_ORANGE, PALE_ORANGE, 18),
    # 物流・輸送
    ("海上運賃\n(スポット)", "4472C4",  PALE_BLUE,  WHITE,       18),
    ("国内輸送\nコスト",    "4472C4",  PALE_BLUE,  WHITE,       18),
    ("リードタイム\n変動",  "4472C4",  PALE_BLUE,  WHITE,       18),
    # 為替・経済
    ("円ドル\n為替",        "70AD47",  PALE_GREEN, PALE_GREEN,  14),
    ("輸入コスト\n影響額",  "70AD47",  PALE_GREEN, PALE_GREEN,  18),
    # 市場・規制
    ("国内市場\n需要動向",  "4472C4",  PALE_BLUE,  WHITE,       18),
    ("環境規制\n動向",      "4472C4",  PALE_BLUE,  WHITE,       18),
    ("競合他社\n動向",      "4472C4",  PALE_BLUE,  WHITE,       18),
    # 取引先・対応
    ("顧客・取引先\nへの影響","375623", PALE_GREEN, PALE_GREEN,  22),
    ("自社在庫\n状況",      "375623",  PALE_GREEN, PALE_GREEN,  16),
    ("対応策\n・優先行動",  "375623",  PALE_GREEN, PALE_GREEN,  24),
    # 管理
    ("緊急度\n高/中/低",   "7030A0",  "EAD1F0",   "EAD1F0",    14),
    ("情報源\nURL/媒体名", DARK_BLUE,  GRAY_FILL,  WHITE,       22),
    ("収集日",             DARK_BLUE,  GRAY_FILL,  WHITE,       13),
    ("備考・\nメモ",       DARK_BLUE,  GRAY_FILL,  WHITE,       28),
]

ROW_TITLE   = 1
ROW_PERIOD  = 2
ROW_HEADER  = 3
ROW_DATA_START = 4

# ── タイトル行 ────────────────────────────────────────────────
ws.merge_cells(f"A{ROW_TITLE}:{get_column_letter(len(columns))}{ROW_TITLE}")
c = ws.cell(ROW_TITLE, 1)
c.value = "食品業界向け包装資材　商品群別ニュース収集フォーマット"
c.fill = fill(DARK_BLUE)
c.font = Font(name="Meiryo", bold=True, color=WHITE, size=14)
c.alignment = align()
ws.row_dimensions[ROW_TITLE].height = 32

# ── 期間・担当行 ─────────────────────────────────────────────
ws.merge_cells(f"A{ROW_PERIOD}:{get_column_letter(len(columns))}{ROW_PERIOD}")
c = ws.cell(ROW_PERIOD, 1)
c.value = "収集期間：　　　　年　　月　　日 ～ 　　　　年　　月　　日　　　担当者名：　　　　　　　　　　　　注目テーマ：中東・ホルムズ海峡情勢"
c.fill = fill(LIGHT_BLUE)
c.font = Font(name="Meiryo", size=10, color=DARK_BLUE, bold=True)
c.alignment = align(h="left")
ws.row_dimensions[ROW_PERIOD].height = 20

# ── ヘッダ行 ─────────────────────────────────────────────────
for col_i, (name, hdr_color, _, _, width) in enumerate(columns, 1):
    c = ws.cell(ROW_HEADER, col_i)
    c.value = name
    c.fill = fill(hdr_color)
    c.font = font(WHITE, bold=True, size=9)
    c.alignment = align()
    c.border = border()
    ws.column_dimensions[get_column_letter(col_i)].width = width
ws.row_dimensions[ROW_HEADER].height = 42

# ── データ行 ─────────────────────────────────────────────────
for row_i, (prod_name, raw_material) in enumerate(products):
    row_num = ROW_DATA_START + row_i
    is_even = (row_i % 2 == 0)

    for col_i, (name, hdr_color, even_color, odd_color, _) in enumerate(columns, 1):
        c = ws.cell(row_num, col_i)
        bg = even_color if is_even else odd_color

        if col_i == 1:          # 商品群列
            c.value = prod_name
            c.fill = fill(MID_BLUE)
            c.font = font(WHITE, bold=True, size=10)
            c.alignment = align()
        elif col_i == 2:        # 主要原材料列
            c.value = raw_material
            c.fill = fill(LIGHT_BLUE)
            c.font = Font(name="Meiryo", size=9, color=DARK_BLUE, bold=True)
            c.alignment = align()
        else:
            c.fill = fill(bg)
            c.font = Font(name="Meiryo", size=9, color="1A1A1A")
            c.alignment = align(h="left")

        c.border = border()

    ws.row_dimensions[row_num].height = 72

# ── ペイン固定（商品群・原材料を固定） ───────────────────────
ws.freeze_panes = "C4"

# ── 印刷設定 ─────────────────────────────────────────────────
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.print_title_rows = f"{ROW_TITLE}:{ROW_HEADER}"


# ════════════════════════════════════════════════════════════════
# シート②：ホルムズ影響度マトリクス（商品群×影響カテゴリ）
# ════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("ホルムズ影響度評価")

impact_columns = [
    "商品群",
    "主要原材料",
    "石油依存度\n(高/中/低)",
    "中東調達\n比率(%)",
    "在庫\n月数",
    "価格転嫁\n難易度\n(高/中/低)",
    "代替調達\n可否",
    "緊急仕入\n必要性\n(高/中/低)",
    "顧客への\n影響度\n(高/中/低)",
    "現状\nステータス",
    "具体的\nリスク内容",
    "対応\nアクション",
    "期限",
    "責任者",
]

# ヘッダ
ws2.merge_cells(f"A1:{get_column_letter(len(impact_columns))}1")
c = ws2.cell(1, 1)
c.value = "ホルムズ海峡封鎖リスク　商品群別影響度評価マトリクス"
c.fill = fill(ORANGE_HDR)
c.font = Font(name="Meiryo", bold=True, color=WHITE, size=13)
c.alignment = align()
ws2.row_dimensions[1].height = 30

impact_col_widths = [18, 20, 16, 14, 12, 16, 14, 16, 16, 20, 30, 30, 13, 13]
impact_col_colors = [
    (DARK_BLUE, MID_BLUE),     # 商品群
    (DARK_BLUE, LIGHT_BLUE),   # 主要原材料
    (ORANGE_HDR, PALE_ORANGE), # 石油依存度
    (ORANGE_HDR, PALE_ORANGE), # 中東調達比率
    ("4472C4", PALE_BLUE),     # 在庫月数
    (ORANGE_HDR, PALE_ORANGE), # 価格転嫁難易度
    ("4472C4", PALE_BLUE),     # 代替調達可否
    (ORANGE_HDR, PALE_ORANGE), # 緊急仕入必要性
    (ORANGE_HDR, PALE_ORANGE), # 顧客影響度
    ("375623", PALE_GREEN),    # 現状ステータス
    ("375623", PALE_GREEN),    # 具体的リスク内容
    ("375623", PALE_GREEN),    # 対応アクション
    (DARK_BLUE, GRAY_FILL),    # 期限
    (DARK_BLUE, GRAY_FILL),    # 責任者
]

for col_i, (name, (hdr_c, _)) in enumerate(zip(impact_columns, impact_col_colors), 1):
    c = ws2.cell(2, col_i)
    c.value = name
    c.fill = fill(hdr_c)
    c.font = font(WHITE, bold=True, size=9)
    c.alignment = align()
    c.border = border()
    ws2.column_dimensions[get_column_letter(col_i)].width = impact_col_widths[col_i - 1]
ws2.row_dimensions[2].height = 50

# 石油系商品に初期値を設定（参考）
oil_dep = {
    "段ボール":               ("低", "5%未満"),
    "発泡スチロール":          ("高", "60〜70%"),
    "食品トレー":              ("高", "60〜70%"),
    "ボードン袋":              ("高", "60〜70%"),
    "ストレッチフィルム":      ("高", "60〜70%"),
    "ニトリル手袋":            ("高", "70〜80%"),
    "プラスチックコンテナー":  ("高", "60〜70%"),
}

for row_i, (prod_name, raw_material) in enumerate(products):
    row_num = 3 + row_i
    is_even = (row_i % 2 == 0)

    dep, ratio = oil_dep.get(prod_name, ("中", "不明"))

    preset = [prod_name, raw_material, dep, ratio, "", "", "", "", "", "", "", "", "", ""]

    for col_i, (val, (hdr_c, data_c)) in enumerate(zip(preset, impact_col_colors), 1):
        c = ws2.cell(row_num, col_i)
        c.value = val
        bg = data_c if not is_even else data_c
        if col_i == 1:
            c.fill = fill(MID_BLUE)
            c.font = font(WHITE, bold=True, size=10)
        elif col_i == 2:
            c.fill = fill(LIGHT_BLUE)
            c.font = Font(name="Meiryo", size=9, color=DARK_BLUE, bold=True)
        elif val == "高":
            c.fill = fill("FFD9D9")
            c.font = Font(name="Meiryo", size=9, color="C00000", bold=True)
        elif val == "低":
            c.fill = fill(PALE_GREEN)
            c.font = Font(name="Meiryo", size=9, color="375623")
        else:
            c.fill = fill(data_c)
            c.font = Font(name="Meiryo", size=9, color="1A1A1A")
        c.alignment = align()
        c.border = border()
    ws2.row_dimensions[row_num].height = 55

ws2.freeze_panes = "C3"
ws2.page_setup.orientation = "landscape"
ws2.page_setup.fitToPage = True
ws2.page_setup.fitToWidth = 1


# ════════════════════════════════════════════════════════════════
# シート③：凡例・使い方ガイド
# ════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("使い方ガイド")

guide = [
    ("【使い方ガイド】", DARK_BLUE, WHITE, True, 14),
    ("", WHITE, "1A1A1A", False, 10),
    ("■ シート①「ニュース収集フォーマット」の使い方", LIGHT_BLUE, DARK_BLUE, True, 11),
    ("① ニュース・業界紙・商社からの情報を収集し、該当商品群の行に記入します。", WHITE, "1A1A1A", False, 10),
    ("② 「緊急度」列は「高」「中」「低」で記入してください。", WHITE, "1A1A1A", False, 10),
    ("③ 「情報源」には、紙名・URL・取引先名などを記入します。", WHITE, "1A1A1A", False, 10),
    ("④ 「収集日」は記入当日の日付を入力してください。", WHITE, "1A1A1A", False, 10),
    ("⑤ オレンジ色の列（中東・ホルムズ関連）は現在最優先で収集してください。", PALE_ORANGE, "C55A11", True, 10),
    ("", WHITE, "1A1A1A", False, 10),
    ("■ シート②「ホルムズ影響度評価」の使い方", LIGHT_BLUE, DARK_BLUE, True, 11),
    ("① 各商品群について、石油依存度・緊急仕入必要性・顧客影響度を「高/中/低」で評価します。", WHITE, "1A1A1A", False, 10),
    ("② 「対応アクション」列に具体的な行動（仕入先交渉、在庫積み増し等）を記入します。", WHITE, "1A1A1A", False, 10),
    ("③ 「期限」と「責任者」を設定し、PDCAを回してください。", WHITE, "1A1A1A", False, 10),
    ("", WHITE, "1A1A1A", False, 10),
    ("■ 主な情報収集先（推奨）", LIGHT_BLUE, DARK_BLUE, True, 11),
    ("・日本包装技術協会（JPI）/ 包装タイムス", WHITE, "1A1A1A", False, 10),
    ("・化学工業日報 / 日刊工業新聞", WHITE, "1A1A1A", False, 10),
    ("・石油化学工業協会 / METI（経済産業省）", WHITE, "1A1A1A", False, 10),
    ("・各原材料メーカー（住友化学、LG化学、旭化成等）の価格レター", WHITE, "1A1A1A", False, 10),
    ("・Bloomberg / Reuters（原油・ナフサ価格、為替）", WHITE, "1A1A1A", False, 10),
    ("・外務省 海外安全情報（中東情勢）", WHITE, "1A1A1A", False, 10),
    ("", WHITE, "1A1A1A", False, 10),
    ("■ 商品別 主要原材料と石油依存度", LIGHT_BLUE, DARK_BLUE, True, 11),
    ("段ボール         ：古紙・クラフト紙（石油依存度：低）　※ただし輸送コストは影響大", WHITE, "1A1A1A", False, 10),
    ("発泡スチロール   ：スチレンモノマー → ベンゼン → ナフサ（石油依存度：高）", PALE_ORANGE, "C55A11", False, 10),
    ("食品トレー       ：PS樹脂・PP樹脂 → ナフサ（石油依存度：高）", PALE_ORANGE, "C55A11", False, 10),
    ("ボードン袋       ：PE樹脂 → エチレン → ナフサ（石油依存度：高）", PALE_ORANGE, "C55A11", False, 10),
    ("ストレッチフィルム：LLDPE樹脂 → エチレン → ナフサ（石油依存度：高）", PALE_ORANGE, "C55A11", False, 10),
    ("ニトリル手袋     ：アクリロニトリル＋ブタジエン → ナフサ（石油依存度：非常に高）", PALE_ORANGE, "C55A11", False, 10),
    ("プラスチックコンテナー：PP・HDPE樹脂 → ナフサ（石油依存度：高）", PALE_ORANGE, "C55A11", False, 10),
]

ws3.column_dimensions["A"].width = 90
for row_i, (text, bg, fg, bold_, size_) in enumerate(guide, 1):
    c = ws3.cell(row_i, 1)
    c.value = text
    c.fill = fill(bg)
    c.font = Font(name="Meiryo", bold=bold_, color=fg, size=size_)
    c.alignment = align(h="left")
    c.border = border()
    ws3.row_dimensions[row_i].height = 20 if text else 8


# ── 保存 ───────────────────────────────────────────────────────
output_path = "/home/user/claudecode/高橋将史商会_商品群別ニュース収集フォーマット.xlsx"
wb.save(output_path)
print(f"✅ 保存完了: {output_path}")
