# -*- coding: utf-8 -*-
# 上期(10〜3月)の営業所別ファイルを担当者コード別に読み取る
import openpyxl, re, json
S="/tmp/claude-0/-home-user-claudecode/7f1b344f-da50-502a-9e34-aab355afdd38/scratchpad/"
FILES={"石見":"iwami","下関":"shimonoseki","松江":"matsue","広域":"kouiki","境港":"sakai","水産部":"suisan"}
FW=str.maketrans("０１２３４５６７８９．","0123456789.")
def sec_key(t):
    t=str(t).translate(FW)
    m=re.match(r"^\s*(\d+)\s*[.,]\s*(.*)$",t)
    if not m: return None
    name=m.group(2)
    if "2023年" in name: return "y2"      # 令和5年10月〜令和6年3月
    if "2024年" in name: return "y1"      # 令和6年10月〜令和7年3月
    if name.startswith("計画") and "比" not in name: return "plan"
    if name.startswith("実績"): return "act"
    return "other"
def num(v): return float(v) if isinstance(v,(int,float)) else 0.0
def parse(office):
    wb=openpyxl.load_workbook(S+f"h1/{FILES[office]}.xlsx",data_only=True); res={}
    for sh in ("全体","機械"):
        ws=wb[sh]; cur=None; d={}
        for r in range(1,ws.max_row+1):
            c=ws.cell(r,3).value
            if isinstance(c,str) and sec_key(c) is not None:
                cur=sec_key(c); d.setdefault(cur,{}); continue
            if cur in (None,"other") or c is None: continue
            lab=str(c).strip(); m=re.match(r"(\d{4})(\d{4})?",lab)
            key=m.group(1) if m else lab
            row=(sum(num(ws.cell(r,x).value) for x in range(5,11)),sum(num(ws.cell(r,x).value) for x in range(12,18)))
            if m and m.group(2): d[cur][key+"+"+m.group(2)]=row
            elif key in d[cur] and not m: continue
            else: d[cur][key]=row
        res[sh]=d
    if office=="広域":
        ws=wb["キョウワ"]; cur=None; d={}
        for r in range(1,ws.max_row+1):
            c=ws.cell(r,3).value
            if isinstance(c,str) and sec_key(c) is not None: cur=sec_key(c); d.setdefault(cur,{}); continue
            if cur in (None,"other") or c is None: continue
            if "税抜" in str(c): d[cur]["税抜"]=(sum(num(ws.cell(r,x).value) for x in range(5,11)),sum(num(ws.cell(r,x).value) for x in range(12,18)))
        res["キョウワ"]=d
    return res
if __name__=="__main__":
    K=openpyxl.load_workbook(S+"sep/kaigi_2603.xlsx",data_only=True)["累計実績"]
    kg={"石見":25,"下関":48,"松江":65,"広域":82,"境港":109,"水産部":117}
    out={}
    for o in FILES:
        p=parse(o); out[o]=p
        print("==",o)
        for sec in ("y2","y1","plan","act"):
            d=p["全体"].get(sec,{}); tot=d.get("合計")
            print(f"  {sec:5} 合計(sales,gp)",[round(x) for x in tot] if tot else None,"codes:",[k for k in d if k!="合計"])
        r=kg[o]; print("  3月kaigi 累計 office:",[K.cell(r,c).value for c in (90,91,92,93)],[K.cell(r,c).value for c in (97,98,99,100)])
    json.dump(out,open(S+"h1_parsed.json","w"),ensure_ascii=False,indent=1)
