# -*- coding: utf-8 -*-
# 9月: 営業所別担当者別実績(上書き) / 営業会議資料 / 第51期9月実績
import openpyxl, re, copy, json, os
from openpyxl.formula.translate import Translator
S="/tmp/claude-0/-home-user-claudecode/7f1b344f-da50-502a-9e34-aab355afdd38/scratchpad/"
SEP=S+"sep/"; W=S+"work/"; CALC=W+"calc/"; OUT=S+"out/"
M=5; MC=5+M; GC=12+M; RC=19+M   # Sep: 値列 J(10), 粗利 Q(17), 率 X(24)
files={"石見":"第51期下期石見営業所担当者別実績.xlsx","下関":"第51期下期下関営業所担当者別実績.xlsx","松江":"第51期下期松江営業所担当者別実績.xlsx","広域":"第51期下期広域営業部担当者別実績.xlsx","境港":"第51期下期境港営業所担当者別実績.xlsx","水産部":"第51期下期水産部部門別実績.xlsx"}
flags=[]
def load_raw(f):
    ws=openpyxl.load_workbook(SEP+f,data_only=True).worksheets[0]; d={}
    for row in ws.iter_rows(min_row=2,values_only=True):
        if row[0] is None: continue
        d[str(row[0]).zfill(4)]={"sales":[float(x or 0) for x in row[2:8]],"gp":[float(x or 0) for x in row[8:14]]}
    return d
RAW=load_raw("2609担当者別売上粗利益実績.xlsx"); RAWM=load_raw("2609担当者別売上粗利益実績機械.xlsx")
def raw(code,kind,m=M,src=None):
    src=src or RAW
    return src.get(code,{"sales":[0]*6,"gp":[0]*6})[kind][m]
wc=openpyxl.load_workbook(SEP+"2609担当者別得意先別実績.xlsx",data_only=True).active
ringer={"sales":0.0,"gp":0.0}
for r in range(2,wc.max_row+1):
    n=str(wc.cell(r,2).value or "")
    if "ﾘﾝｶﾞｰ" in n or "リンガー" in n:
        ringer["sales"]+=(wc.cell(r,11).value or 0)/1000; ringer["gp"]+=(wc.cell(r,12).value or 0)/1000
ringer_m={"sales":0.0,"gp":0.0}
wd=openpyxl.load_workbook(SEP+"2609タカハシ包装売上全明細.xlsx",read_only=True,data_only=True).worksheets[0]
for row in wd.iter_rows(min_row=2,values_only=True):
    n=str(row[2] or "")
    if row[0] is not None and str(row[14])=="08" and ("ﾘﾝｶﾞｰ" in n or "リンガー" in n):
        ringer_m["sales"]+=float(row[10] or 0)/1000; ringer_m["gp"]+=float(row[13] or 0)/1000
print("Ringer Sep:",ringer,"machine:",ringer_m)
SEC_KEYS={"1":"y2024","2":"y2025","4":"plan","7":"act"}
def sec_start(ws):
    secs={}
    for r in range(1,ws.max_row+1):
        c=ws.cell(r,3).value
        if isinstance(c,str) and len(c)>=2 and c[1] in "．,":
            secs[c[0].translate(str.maketrans("１２３４５６７８９","123456789"))]=r
    return secs
def num(v): return float(v) if isinstance(v,(int,float)) else 0.0
def parse_office(o):
    wb=openpyxl.load_workbook(CALC+files[o],data_only=True); res={}
    for sh in ["全体","機械"]:
        ws=wb[sh]; secs=sec_start(ws); res[sh]={}
        order=sorted(secs.items(),key=lambda x:x[1])
        for i,(k,st) in enumerate(order):
            if k not in SEC_KEYS: continue
            en=order[i+1][1] if i+1<len(order) else ws.max_row+1
            d={}
            for r in range(st+1,en):
                c=ws.cell(r,3).value
                if c is None: continue
                lab=str(c).strip(); m=re.match(r"(\d{4})",lab); key=m.group(1) if m else lab
                d[key]={"sales":[num(ws.cell(r,5+j).value) for j in range(6)],"gp":[num(ws.cell(r,12+j).value) for j in range(6)],"row":r}
            res[sh][SEC_KEYS[k]]=d
    return res
OFFICE={o:parse_office(o) for o in files}
# ---- 8月実績(ファイル) vs 最新RAW 8月列 の差異チェック
for o in files:
    for sh,src in [("全体",RAW),("機械",RAWM)]:
        for code,v in OFFICE[o][sh]["act"].items():
            if not re.match(r"\d{4}",code): continue
            fs,fg=v["sales"][M-1],v["gp"][M-1]; rs,rg=raw(code,"sales",M-1,src),raw(code,"gp",M-1,src)
            if abs(fs-rs)>0.01 or abs(fg-rg)>0.01: flags.append(f"8月差異 {o} {sh} {code}: ファイル {fs:.3f}/{fg:.3f} vs 最新RAW {rs:.3f}/{rg:.3f}")
wk=openpyxl.load_workbook(CALC+files["広域"],data_only=True)["キョウワ"]
def kyowa_rows():
    secs=sec_start(wk); order=sorted(secs.items(),key=lambda x:x[1]); res={}
    for i,(k,st) in enumerate(order):
        en=order[i+1][1] if i+1<len(order) else wk.max_row+1
        for r in range(st+1,en):
            c=wk.cell(r,3).value
            if c and "税抜" in str(c):
                res[k]={"sales":[num(wk.cell(r,5+j).value) for j in range(6)],"gp":[num(wk.cell(r,12+j).value) for j in range(6)],"row":r}
    return res
KY=kyowa_rows(); print("Kyowa secs:",{k:(v['row'],v['sales'][M]) for k,v in KY.items()})
def kyowa(key,kind,m=M):
    k={"y2024":"1","y2025":"2","plan":"4","act":"7"}[key]
    if key=="act": return round(raw("0502",kind,m))
    return KY[k][kind][m]
def act(o,sh,code,kind,m):
    if m==M: return raw(code,kind,m,RAWM if sh=="機械" else RAW)
    return OFFICE[o][sh]["act"].get(code,{"sales":[0]*6,"gp":[0]*6})[kind][m]
def hist(o,sh,key,code,kind,m): return OFFICE[o][sh][key].get(code,{"sales":[0]*6,"gp":[0]*6})[kind][m]
def codes_in(o,sh="全体"): return [k for k in OFFICE[o][sh]["act"] if re.match(r"\d{4}",k)]
def office_total(o,sh,key,kind,m):
    if key=="act":
        if m==M: return sum(act(o,sh,c,kind,m) for c in codes_in(o,sh))
        return OFFICE[o][sh]["act"]["合計"][kind][m]
    return OFFICE[o][sh][key]["合計"][kind][m]
# ======================= 1. 営業所別ファイル 9月列
def translate_fill(ws,r,cols=(MC,GC,RC)):
    for c in cols:
        cell=ws.cell(r,c); left=ws.cell(r,c-1)
        if cell.value is None and isinstance(left.value,str) and left.value.startswith("="):
            cell.value=Translator(left.value,origin=left.coordinate).translate_formula(cell.coordinate)
            cell.number_format=left.number_format; cell._style=copy.copy(left._style)
def update_office(o):
    wb=openpyxl.load_workbook(W+files[o])
    for sh in ["全体","機械","包装資材"]:
        ws=wb[sh]; secs=sec_start(ws); order=sorted(secs.items(),key=lambda x:x[1])
        for i,(k,st) in enumerate(order):
            en=order[i+1][1] if i+1<len(order) else ws.max_row+1
            if k not in ("7","8","9","5","6"): continue
            for r in range(st+1,en):
                c=ws.cell(r,3).value
                if c is None: continue
                lab=str(c).strip(); m=re.match(r"(\d{4})",lab)
                if k=="7" and sh in ("全体","機械") and m and not isinstance(ws.cell(r,MC-1).value,str):
                    code=m.group(1); src=RAWM if sh=="機械" else RAW
                    ws.cell(r,MC).value=raw(code,"sales",M,src); ws.cell(r,GC).value=raw(code,"gp",M,src)
                    for cc in (MC,GC): ws.cell(r,cc).number_format=ws.cell(r,cc-1).number_format
                if k=="7" and lab=="内リンガーハット":
                    if sh=="全体": ws.cell(r,MC).value=round(ringer["sales"],3); ws.cell(r,GC).value=round(ringer["gp"],3)
                    elif sh=="機械": ws.cell(r,MC).value=round(ringer_m["sales"],3); ws.cell(r,GC).value=round(ringer_m["gp"],3)
                if k=="7" and lab.startswith("【参考】大海分"): continue
                translate_fill(ws,r)
    if o=="広域":
        ws=wb["キョウワ"]; secs=sec_start(ws); order=sorted(secs.items(),key=lambda x:x[1])
        for i,(k,st) in enumerate(order):
            en=order[i+1][1] if i+1<len(order) else ws.max_row+1
            for r in range(st+1,en):
                c=ws.cell(r,3).value
                if c and "税抜" in str(c):
                    if k=="7": ws.cell(r,MC).value=round(raw("0502","sales")); ws.cell(r,GC).value=round(raw("0502","gp"))
                    translate_fill(ws,r)
    wb.save(OUT+files[o]); print("saved office",o)
for o in files: update_office(o)
# ======================= 2. 営業会議資料 9月
NAME2CODE={("石見","佐々木課長"):"0011",("石見","川上"):"0012",("石見","三浦"):"0013",("石見","木村課長"):"0014",("石見","その他"):"0019",("石見","自治体"):"0020",
 ("下関","高橋所長"):"0021",("下関","東野"):"0024",("下関","橋本課長"):"0025",("下関","井上課長"):"0026",("下関","中国"):"0027",("下関","自治体"):"0028",("下関","その他"):"0029",
 ("松江","飯塚主任"):"0032",("松江","前田課長代理"):"0033",("松江","原田"):"0034",("松江","自治体"):"0035",("松江","その他"):"0039",
 ("広域","馬場部長（鳥栖）"):"0042",("広域","ペーパーハグ"):"0045",("広域","角田"):"0046",("広域","馬場部長（静岡）"):"0047",("広域","馬場所長（東京）"):"0550",("広域","馬場部長（東京）"):"0552",
 ("境港","西岡"):"0903",("境港","足立次長"):"0904",("境港","藤波課長"):"0909",("境港","足立課長"):"0910",("境港","川邊主任"):"0911",("境港","白根主任"):"0914",("境港","海外"):"0918",("境港","その他"):"0919",("境港","自治体"):"0930"}
OFFICE_ROWS={"石見":(7,29),"下関":(30,55),"松江":(56,75),"広域":(76,99),"境港":(100,131)}
def kaigi_values(o,code,sh,months):
    v={}
    for col,key in [("F","y2024"),("G","y2025"),("H","plan")]: v[col]=sum(hist(o,sh,key,code,"sales",m) for m in months)
    v["I"]=sum(act(o,sh,code,"sales",m) for m in months)
    for col,key in [("M","y2024"),("N","y2025"),("O","plan")]: v[col]=sum(hist(o,sh,key,code,"gp",m) for m in months)
    v["P"]=sum(act(o,sh,code,"gp",m) for m in months)
    return v
def total_values(o,sh,months):
    v={}
    for col,key in [("F","y2024"),("G","y2025"),("H","plan")]: v[col]=sum(office_total(o,sh,key,"sales",m) for m in months)
    v["I"]=sum(office_total(o,sh,"act","sales",m) for m in months)
    for col,key in [("M","y2024"),("N","y2025"),("O","plan")]: v[col]=sum(office_total(o,sh,key,"gp",m) for m in months)
    v["P"]=sum(office_total(o,sh,"act","gp",m) for m in months)
    return v
def fill_kaigi_sheet(ws,months,label_m,label_h,check=False):
    if not check:
        for col in ["F","M","T"]: ws[f"{col}5"]=f"令和6年{label_m}"
        for col in ["G","N","U"]: ws[f"{col}5"]=f"令和7年{label_m}"
        for col in ["H","I","O","P","V","W"]: ws[f"{col}5"]=f"令和8年{label_m}"
        ws["A2"]=label_h
    cur=None; pend=None
    for r in range(7,144):
        d=ws.cell(r,4).value; e=ws.cell(r,5).value; b=ws.cell(r,2).value; a=ws.cell(r,1).value
        for o,(r1,r2) in OFFICE_ROWS.items():
            if r1<=r<=r2: cur=o
        def put(v):
            for col,val in v.items():
                if check:
                    old=ws[f"{col}{r}"].value
                    if not isinstance(old,(int,float)) or abs(old-val)>0.5: flags.append(f"8月単月再現チェック 営業会議 {col}{r} ({d or e or b or a}): 既存 {old} vs 再計算 {val:.3f}")
                else: ws[f"{col}{r}"]=val
        if d and cur and (cur,d) in NAME2CODE:
            code=NAME2CODE[(cur,d)]; put(kaigi_values(cur,code,"全体",months)); pend=(cur,code)
        elif e=="機械" and isinstance(ws.cell(r,6).value,(int,float)) and cur and pend:
            put(kaigi_values(pend[0],pend[1],"機械",months))
        elif b=="水産部": put(total_values("水産部","全体",months)); pend=None
        elif e=="荷役" and r==137:
            v1=kaigi_values("水産部","0104","全体",months); v2=kaigi_values("水産部","0106","全体",months); put({k:v1[k]+v2[k] for k in v1})
        elif e=="機械" and r==138: put(total_values("水産部","機械",months))
        elif a and "キョウワ" in str(a):
            put({"F":sum(kyowa("y2024","sales",m) for m in months),"G":sum(kyowa("y2025","sales",m) for m in months),"H":sum(kyowa("plan","sales",m) for m in months),"I":sum(kyowa("act","sales",m) for m in months),
                 "M":sum(kyowa("y2024","gp",m) for m in months),"N":sum(kyowa("y2025","gp",m) for m in months),"O":sum(kyowa("plan","gp",m) for m in months),"P":sum(kyowa("act","gp",m) for m in months)})
        elif d and cur and d not in ("営業","自治体","中国事業","海外","【参考】馬場部長合計") and not check:
            flags.append(f"営業会議資料 row {r}: 未対応ラベル {d}")
def make_kaigi():
    wb=openpyxl.load_workbook(W+"【営業会議資料】令和8年8月実績.xlsx")
    fill_kaigi_sheet(wb["単月"],[M-1],"","",check=True)   # 8月の既存値を再現できるか検証
    fill_kaigi_sheet(wb["単月"],[M],"9月","1．単月実績")
    fill_kaigi_sheet(wb["下期累計"],list(range(0,M+1)),"4〜9月","2．下期累計実績（4〜9月）")
    wb.save(OUT+"【営業会議資料】令和8年9月実績.xlsx"); print("saved kaigi")
make_kaigi()
# ======================= 3. 第51期9月実績
def make_jisseki():
    wb=openpyxl.load_workbook(W+"第51期8月実績.xlsx")
    ws8=wb["8月"]; bd=wb.copy_worksheet(ws8); bd.title="BACKDATA8月"
    wb._sheets.remove(bd); wb._sheets.insert(wb.sheetnames.index("期初来累計")+1,bd)
    ws8.title="9月"; ws=ws8
    for col,t in zip("IJKL",["2024年9月売上実績","2025年9月売上実績","2026年9月売上計画","2026年9月売上実績"]): ws[f"{col}3"]=t
    for col,t in zip("PQRS",["2024年9月粗利益実績","2025年9月粗利益実績","2026年9月粗利益計画","2026年9月粗利益実績"]): ws[f"{col}3"]=t
    for col,t in zip("WXYZ",["2024年9月粗利益率実績","2025年9月粗利益率実績","2026年9月粗利益率計画","2026年9月粗利益率実績"]): ws[f"{col}3"]=t
    def put(r,vals):
        for col,v in zip(["I","J","K","L","P","Q","R","S"],vals): ws[f"{col}{r}"]=v
    def tot(o,sh): return tuple(office_total(o,sh,k,"sales",M) for k in ["y2024","y2025","plan","act"])+tuple(office_total(o,sh,k,"gp",M) for k in ["y2024","y2025","plan","act"])
    def per(o,code,sh="全体"): return tuple(hist(o,sh,k,code,"sales",M) for k in ["y2024","y2025","plan"])+(act(o,sh,code,"sales",M),)+tuple(hist(o,sh,k,code,"gp",M) for k in ["y2024","y2025","plan"])+(act(o,sh,code,"gp",M),)
    put(5,tot("石見","全体")); put(8,tot("石見","機械")); put(9,per("石見","0020"))
    put(10,tot("下関","全体")); put(13,tot("下関","機械")); put(14,per("下関","0028")); put(15,per("下関","0027"))
    put(16,tot("松江","全体")); put(19,tot("松江","機械")); put(20,per("松江","0035"))
    put(21,tot("広域","全体")); put(23,tot("広域","機械"))
    rg=OFFICE["広域"]["全体"]; rgm=OFFICE["広域"]["機械"]
    def rl(d,key):
        for k in d[key]:
            if "リンガー" in k and "以外" not in k: return d[key][k]
    put(24,(rl(rg,"y2024")["sales"][M],rl(rg,"y2025")["sales"][M],rl(rg,"plan")["sales"][M],round(ringer["sales"],3),rl(rg,"y2024")["gp"][M],rl(rg,"y2025")["gp"][M],rl(rg,"plan")["gp"][M],round(ringer["gp"],3)))
    put(26,(rl(rgm,"y2024")["sales"][M],rl(rgm,"y2025")["sales"][M],rl(rgm,"plan")["sales"][M],round(ringer_m["sales"],3),rl(rgm,"y2024")["gp"][M],rl(rgm,"y2025")["gp"][M],rl(rgm,"plan")["gp"][M],round(ringer_m["gp"],3)))
    put(30,tot("境港","全体")); put(33,tot("境港","機械")); put(34,per("境港","0930")); put(35,per("境港","0918"))
    put(41,tot("水産部","全体"))
    v1=per("水産部","0104"); v2=per("水産部","0106"); put(43,tuple(a+b for a,b in zip(v1,v2)))
    put(44,tot("水産部","機械"))
    put(50,(kyowa("y2024","sales"),kyowa("y2025","sales"),kyowa("plan","sales"),kyowa("act","sales"),kyowa("y2024","gp"),kyowa("y2025","gp"),kyowa("plan","gp"),kyowa("act","gp")))
    wsd=wb["下期"]; n=0
    for row in wsd.iter_rows():
        for c in row:
            if isinstance(c.value,str) and "'8月'!" in c.value:
                c.value=re.sub(r"'8月'!(\$?[A-Z]+\$?\d+)",r"'9月'!\1+BACKDATA8月!\1",c.value); n+=1
    print("下期 formulas updated:",n)
    for w in wb.worksheets:
        if w.title=="下期": continue
        for row in w.iter_rows():
            for c in row:
                if isinstance(c.value,str) and "'8月'!" in c.value: flags.append(f"第51期9月実績 {w.title}!{c.coordinate} still references '8月'")
    wb.save(OUT+"第51期9月実績.xlsx"); print("saved jisseki")
make_jisseki()
json.dump({"ringer":ringer,"ringer_m":ringer_m},open(OUT+"ringer.json","w"),ensure_ascii=False)
json.dump(flags,open(OUT+"flags.json","w"),ensure_ascii=False,indent=1)
print("FLAGS:",*flags,sep="\n ")
