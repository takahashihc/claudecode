# 第51期実績: 階層を 営業所＝営業＋自治体(＋海外)、営業＝包装資材＋機械 に組み替え（値は不変）
# 入力(リーフ)は 包装資材・機械・自治体・海外。広域は 営業部＝包装資材＋機械、水産部は 包装資材＋荷役＋機械。
import openpyxl, sys
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter as L
src,calc,dst=sys.argv[1],sys.argv[2],sys.argv[3]   # calc = 同ファイルをLibreOfficeで再計算したもの（キャッシュ値取得用）
wb=openpyxl.load_workbook(src); wv=openpyxl.load_workbook(calc,data_only=True)
# (office, eigyo, pack, kikai, extras)
BLOCKS=[(5,6,7,8,[9]),(10,11,12,13,[14,15]),(16,17,18,19,[20]),(30,31,32,33,[34,35])]
def colsets(sheet):
    if sheet=="BACKDATA上期": return [[s+i for i in (0,1,2,3,7,8,9,10)] for s in (9,23,37,51,65,79,93)], 93
    return [[9,10,11,12,16,17,18,19]], None
SUMMODE={"下期","期初来累計"}
log=[]
for ws in wb.worksheets:
    v=wv[ws.title]; groups,cum=colsets(ws.title)
    for cols in groups:
        sum_mode = ws.title in SUMMODE   # 上期は確定値(固定入力)を保つため包装資材は逆算した固定値
        for c in cols:
            cl=L(c)
            def num(r):
                x=v.cell(r,c).value
                return x if isinstance(x,(int,float)) else None
            def num0(r):   # 空欄は0扱い（エラー値は不可）
                x=v.cell(r,c).value
                return 0 if x is None else (x if isinstance(x,(int,float)) else None)
            def setpack(r_pack,r_ref,r_office=None,r_others=()):  # リーフ化
                cell=ws.cell(r_pack,c)
                f=ws.cell(r_ref,c).value
                if sum_mode and isinstance(f,str) and f.startswith("="):
                    cell.value=Translator(f,origin=f"{cl}{r_ref}").translate_formula(f"{cl}{r_pack}")
                else:   # 固定値: 営業所合計（確定値）−機械−自治体−海外 で逆算し、合計値を不変に保つ
                    vals=[num0(r_office)]+[num0(x) for x in r_others]
                    assert all(x is not None for x in vals),(ws.title,cl,r_pack,vals)
                    cell.value=vals[0]-sum(vals[1:])
                log.append((ws.title,cell.coordinate))
            skipped=[]
            for o,e,p,k,ex in BLOCKS:
                if sum_mode is False and any(num0(r) is None for r in [o,k]+ex):
                    skipped.append((ws.title,cl,o)); continue
                setpack(p,k,o,[k]+ex)
                ws.cell(e,c).value=f"={cl}{p}+{cl}{k}"
                ws.cell(o,c).value="="+"+".join(f"{cl}{x}" for x in [e]+ex)
            # 広域: 営業部=包装資材+機械、リンガーハット=包装+機械、以外=差分
            if all(num(r) is not None for r in (21,23,24,26)):   # 上期の月次ブロックは広域が元から#REF!のため対象外
                setpack(22,23,21,[23]); ws.cell(21,c).value=f"={cl}22+{cl}23"
                setpack(25,26,24,[26]); ws.cell(24,c).value=f"={cl}25+{cl}26"
                ws.cell(28,c).value=f"={cl}22-{cl}25"; ws.cell(29,c).value=f"={cl}23-{cl}26"; ws.cell(27,c).value=f"={cl}28+{cl}29"
            # 水産部: 包装資材+荷役+機械
            if sum_mode is False and any(num0(r) is None for r in (41,43,44)):
                skipped.append((ws.title,cl,41)); continue
            setpack(42,43,41,[43,44]); ws.cell(41,c).value=f"={cl}42+{cl}43+{cl}44"
wb.save(dst); print(len(log),"leaf cells set")
