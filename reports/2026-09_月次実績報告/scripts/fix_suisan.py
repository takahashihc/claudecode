# 水産部ファイル: 6〜9セクションの「合計」が境港水産部合計(直上行)でなく JF境(3行上)を参照しているテンプレート不備を修正
import openpyxl, re
O="/tmp/claude-0/-home-user-claudecode/7f1b344f-da50-502a-9e34-aab355afdd38/scratchpad/out/"
wb=openpyxl.load_workbook(O+"第51期下期水産部部門別実績.xlsx"); log=[]
for sh in ["全体","機械","包装資材"]:
    ws=wb[sh]
    for r in range(1,ws.max_row+1):
        if ws.cell(r,3).value!="合計": continue
        if not (str(ws.cell(r-1,3).value or "").startswith("境港水産部合計") and str(ws.cell(r-3,3).value or "").startswith("0031")): continue
        for c in range(5,19):
            v=ws.cell(r,c).value
            m=re.fullmatch(r"=([A-Z]+)(\d+)\+\1(\d+)\+\1(\d+)",str(v))
            if m and int(m.group(4))==r-3:
                col=m.group(1); new=f"={col}{m.group(2)}+{col}{m.group(3)}+{col}{r-1}"
                ws.cell(r,c).value=new; log.append(f"{sh}!{ws.cell(r,c).coordinate}: {v} -> {new}")
wb.save(O+"第51期下期水産部部門別実績.xlsx"); print(len(log),"fixed"); print(*log[:6],sep="\n"); print("...")
