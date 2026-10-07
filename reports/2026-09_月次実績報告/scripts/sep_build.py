import json, copy, sys
import copy as _copy
from pptx.util import Emu
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from lxml import etree
S="/tmp/claude-0/-home-user-claudecode/7f1b344f-da50-502a-9e34-aab355afdd38/scratchpad/"
D=json.load(open(S+"sep_final.json"))
O=D["offices"]; OFF=["石見","下関","松江","広域","境港","水産部","キョウワ"]
NS="{http://schemas.openxmlformats.org/drawingml/2006/main}"
def pct(a,b): return a/b*100
def f0(x): return f"{x:,.0f}"
def p1(x): return f"{x:.1f}%"
def sgn(x): return f"+{x:,.0f}" if x>=0 else f"▲{abs(x):,.0f}"
def get(o,m,p,k): return O[o][m][p][k]
def yoy(o,m,p): return pct(get(o,m,p,"act"),get(o,m,p,"prior"))
def vsp(o,m,p): return pct(get(o,m,p,"act"),get(o,m,p,"plan"))
def rate(o,p,k="act"): return pct(get(o,"gp",p,k),get(o,"sales",p,k))
T="合計"
PV=json.load(open(S+"pricevol2_sep.json")); PVT=PV["タカハシ包装(6)"]; PVI=PV["石見"]
T6="タカハシ包装"   # 6営業所・部（キョウワを含まない）
O[T6]={m:{p:{k:sum(O[o][m][p][k] for o in OFF[:6]) for k in ["act","plan","prior"]} for p in ["m","h2","ytd"]} for m in ["sales","gp"]}
# ---------- key figures
m_gp=get(T,"gp","m","act"); m_sa=get(T,"sales","m","act")
h_gp=get(T,"gp","h2","act"); h_sa=get(T,"sales","h2","act")
y_gp=get(T,"gp","ytd","act"); y_sa=get(T,"sales","ytd","act")
gpdiff={o:get(o,"gp","m","act")-get(o,"gp","m","prior") for o in OFF}
tot_gpdiff=m_gp-get(T,"gp","m","prior")
ex={m:{k:get(T,m,"m",k)-get("境港",m,"m",k) for k in ["act","plan","prior"]} for m in ["sales","gp"]}
ex_sa_yoy=pct(ex["sales"]["act"],ex["sales"]["prior"]); ex_gp_yoy=pct(ex["gp"]["act"],ex["gp"]["prior"]); ex_gp_plan=pct(ex["gp"]["act"],ex["gp"]["plan"])
plus=sorted([o for o in OFF if gpdiff[o]>0],key=lambda o:-gpdiff[o])
minus=sorted([o for o in OFF if gpdiff[o]<0],key=lambda o:gpdiff[o])
def joinc(lst,n=3): return "・".join(f"{o}{sgn(gpdiff[o])}" for o in lst[:n])
print("KEY:",f0(m_gp),p1(yoy(T,'gp','m')),p1(vsp(T,'gp','m')),"| sales",f0(m_sa),p1(yoy(T,'sales','m')),p1(vsp(T,'sales','m')),"| rate",p1(rate(T,'m')),p1(rate(T,'m','prior')),"| ex-sakai",p1(ex_sa_yoy),p1(ex_gp_yoy),p1(ex_gp_plan))
# ---------- helpers
def set_text(shape,text):
    tf=shape.text_frame
    paras=tf.paragraphs
    p=paras[0]
    runs=p.runs
    if runs:
        runs[0].text=text
        for r in runs[1:]: r._r.getparent().remove(r._r)
    else:
        p.text=text
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
def set_cell(cell,text,color=None):
    p=cell.text_frame.paragraphs[0]
    r=p.runs[0]; r.text=text
    for rr in p.runs[1:]: rr._r.getparent().remove(rr._r)
    if color:
        rPr=r._r.get_or_add_rPr()
        sf=rPr.find(NS+"solidFill")
        if sf is not None: rPr.remove(sf)
        sf=etree.SubElement(rPr,NS+"solidFill"); c=etree.SubElement(sf,NS+"srgbClr"); c.set("val",color)
        # solidFill must come before latin/ea: reorder
        for tag in ["latin","ea","cs"]:
            el=rPr.find(NS+tag)
            if el is not None: rPr.remove(el); rPr.append(el)
def ratio_color(v):
    v=round(v,1)
    return "1A6E1A" if v>=100 else ("C55A11" if v>=90 else "C0392B")
def shapes_by_name(slide): return {sh.name:sh for sh in slide.shapes}
def slide_shapes_cache(slide): return list(slide.shapes)
from pptx.util import Pt
prs=Presentation(S+"work/令和8年8月実績報告.pptx")
sl=prs.slides
FOOT="株式会社タカハシ包装センター | 令和8年9月実績"
tr=D["trend"]
# ---------- slide 1 (cover)
s=shapes_by_name(sl[0])
set_text(s["TextBox 3"],"令和8年9月　月次実績報告")
set_text(s["TextBox 4"],"第51期下期 9月実績（第51期通期確定）　～粗利益額を主軸とした分析～")
set_text(s["TextBox 7"],f"粗利益額（キョウワ含む）: 9月 {f0(m_gp)}千円（前期比{p1(yoy(T,'gp','m'))}・計画比{p1(vsp(T,'gp','m'))}）")
set_text(s["TextBox 8"],f"・ 第51期通期は売上{p1(yoy(T,'sales','ytd'))}・粗利{p1(yoy(T,'gp','ytd'))}（計画比{p1(vsp(T,'gp','ytd'))}）で増収増益着地。下期粗利は前期比{p1(yoy(T,'gp','h2'))}")
set_text(s["TextBox 9"],f"・ 9月は境港{sgn(gpdiff['境港'])}（前期機械大口の反動）を {joinc(plus,3)} が吸収し増益")
for nm in ("Rectangle 10","TextBox 11","TextBox 12"):
    el=s[nm]._element; el.getparent().remove(el)
s["TextBox 13"].top=Emu(int(4.35*914400)); s["TextBox 14"].top=Emu(int(4.68*914400))
set_text(s["TextBox 13"],"分析期間: 単月（9月）／下期（4〜9月）／期初来累計（令和7年10月〜令和8年9月＝第51期通期）")
set_text(s["TextBox 14"],"作成日: 令和8年10月7日　　資料作成者: 高橋将史")
# ---------- slide 2 (KPI cards)
s=shapes_by_name(sl[1])
set_text(s["TextBox 5"],FOOT)
cards=[("TextBox 8","TextBox 9","TextBox 11","9月 粗利益額",f0(m_gp),f"前期比 {p1(yoy(T,'gp','m'))} / 計画比 {p1(vsp(T,'gp','m'))}"),
       ("TextBox 14","TextBox 15","TextBox 17","9月 売上高",f0(m_sa),f"前期比 {p1(yoy(T,'sales','m'))} / 計画比 {p1(vsp(T,'sales','m'))}"),
       ("TextBox 20","TextBox 21","TextBox 23","9月 粗利率",f"{rate(T,'m'):.1f}",f"前期 {p1(rate(T,'m','prior'))} → 今期"),
       ("TextBox 26","TextBox 27","TextBox 29","下期 粗利益額",f0(h_gp),f"前期比 {p1(yoy(T,'gp','h2'))} / 計画比 {p1(vsp(T,'gp','h2'))}"),
       ("TextBox 32","TextBox 33","TextBox 35","下期 売上高",f0(h_sa),f"前期比 {p1(yoy(T,'sales','h2'))} / 計画比 {p1(vsp(T,'sales','h2'))}"),
       ("TextBox 38","TextBox 39","TextBox 41","下期 粗利率",f"{rate(T,'h2'):.1f}",f"前期 {p1(rate(T,'h2','prior'))} → 今期"),
       ("TextBox 44","TextBox 45","TextBox 47","通期 粗利益額",f0(y_gp),f"前期比 {p1(yoy(T,'gp','ytd'))} / 計画比 {p1(vsp(T,'gp','ytd'))}"),
       ("TextBox 50","TextBox 51","TextBox 53","通期 売上高",f0(y_sa),f"前期比 {p1(yoy(T,'sales','ytd'))} / 計画比 {p1(vsp(T,'sales','ytd'))}"),
       ("TextBox 56","TextBox 57","TextBox 59","通期 粗利率",f"{rate(T,'ytd'):.1f}",f"前期 {p1(rate(T,'ytd','prior'))} → 今期")]
for a,b,c,t1,t2,t3 in cards:
    set_text(s[a],t1); set_text(s[b],t2); set_text(s[c],t3)
# 各カード下段: タカハシ包装のみ（キョウワ除く）
def tk(m,p): return O[T6][m][p]
sub=[]
for p in ["m","h2","ytd"]:
    g=tk("gp",p); sa=tk("sales",p)
    sub.append(f"タカハシ包装のみ {f0(g['act'])}（前期比{p1(pct(g['act'],g['prior']))}）")
    sub.append(f"タカハシ包装のみ {f0(sa['act'])}（前期比{p1(pct(sa['act'],sa['prior']))}）")
    sub.append(f"タカハシ包装のみ {p1(pct(g['act'],sa['act']))}（前期{p1(pct(g['prior'],sa['prior']))}）")
for (a,b,c,t1,t2,t3),txt in zip(cards,sub):
    src=s[c]; el=_copy.deepcopy(src._element); src._element.getparent().append(el)
    from pptx.shapes.autoshape import Shape
    nb=[x for x in slide_shapes_cache(sl[1]) if x._element is el][0]
    nb.left=s[b].left-Emu(int(0.06*914400)); nb.width=Emu(int(2.72*914400)); nb.top=src.top+Emu(int(0.19*914400)); nb.height=Emu(int(0.2*914400))
    set_text(nb,txt)
    for r in nb.text_frame.paragraphs[0].runs: r.font.size=Pt(7.5)
set_text(s["TextBox 3"],"3期間サマリー（単位: 千円）　※粗利益額を主指標／上段＝タカハシ包装＋キョウワ、下段＝タカハシ包装のみ／通期＝第51期")
# ---------- slides 3-5 (tables)
names={"石見":"石見営業所","下関":"下関営業所","松江":"松江営業所","広域":"広域営業部","境港":"境港営業所","水産部":"水産部（計）","キョウワ":"キョウワ"}
import copy as _copy
from pptx.util import Emu
def fill_table(slide,period):
    gf=[sh for sh in slide.shapes if sh.has_table][0]; t=gf.table; tbl=t._tbl
    if len(tbl.tr_lst)==9:   # 合計行(最終行)の書式を複製し、水産部の次に「タカハシ包装合計」行を追加
        newtr=_copy.deepcopy(tbl.tr_lst[8]); tbl.tr_lst[7].addprevious(newtr)
    if len(tbl.tblGrid.gridCol_lst)==8:   # 前期比・計画比の右に差額列を追加（売上・粗利とも）
        grid=tbl.tblGrid.gridCol_lst
        for src_i in (7,4,3,2):   # 後ろから挿入して添字ずれを避ける: 粗利計画比,粗利前期比... 
            pass
        def dup(tr_i,after_i):
            for tr in tbl.tr_lst:
                tcs=tr.tc_lst; n=_copy.deepcopy(tcs[after_i]); tcs[after_i].addnext(n)
        # 元の列: 0営業所 1売上 2前期比 3計画比 4粗利 5前期比 6計画比 7粗利率
        for after_i in (6,5,3,2):   # 後ろから: 粗利計画比→粗利前期比→売上計画比→売上前期比 の右に複製
            dup(None,after_i)
            g=_copy.deepcopy(grid[after_i]); grid[after_i].addnext(g); grid=tbl.tblGrid.gridCol_lst
        # 差額列ヘッダー（2段目）を書き換え
        hdr=tbl.tr_lst[0].tc_lst
        for ci,txt in [(3,"前期差額"),(5,"計画差額"),(8,"前期差額"),(10,"計画差額")]:
            hdr[ci].txBody.p_lst[1].r_lst[0].t.text=txt if False else txt
    widths=[1.95,0.93,0.58,0.78,0.58,0.78,0.78,0.58,0.78,0.58,0.78,0.55]
    for c,wd in zip(t.columns,widths): c.width=Emu(int(wd*914400))
    gf.left=Emu(int(0.18*914400)); gf.width=Emu(int(sum(widths)*914400))
    for r in t.rows: r.height=Emu(329184)   # 0.36in x10行 = 3.6in
    rows=OFF[:6]+[T6,"キョウワ",T]
    for ri,o in enumerate(rows,start=1):
        total=(o in (T,T6))
        sa=O[o]["sales"][period]; gp=O[o]["gp"][period]
        label={T6:"タカハシ包装合計",T:"タカハシ包装＋キョウワ合計"}.get(o,names.get(o))
        def trio(d):
            return [(pct(d["act"],d["prior"]),d["act"]-d["prior"]),(pct(d["act"],d["plan"]),d["act"]-d["plan"])]
        set_cell(t.cell(ri,0),label); set_cell(t.cell(ri,1),f0(sa["act"])); set_cell(t.cell(ri,6),f0(gp["act"]))
        for base,d in ((2,sa),(7,gp)):
            for k,(rate_,diff) in enumerate(trio(d)):
                col=None if total else ratio_color(rate_)
                set_cell(t.cell(ri,base+2*k),p1(rate_),col); set_cell(t.cell(ri,base+2*k+1),sgn(diff),col)
        set_cell(t.cell(ri,11),p1(pct(gp["act"],sa["act"])))
    for tr in tbl.tr_lst:   # セル余白を詰め、率・差額の文字を9ptに
        for ci,tc in enumerate(tr.tc_lst):
            tc.tcPr.set("marL","22860"); tc.tcPr.set("marR","22860")
            if ci in (2,3,4,5,7,8,9,10):
                for p_ in tc.txBody.p_lst:
                    for r_ in p_.r_lst: r_.get_or_add_rPr().set("sz","900")
s=shapes_by_name(sl[2]); set_text(s["TextBox 5"],FOOT)
set_text(s["TextBox 2"],"営業所別　9月単月実績　～売上・粗利益額～")
fill_table(sl[2],"m")
set_text(s["TextBox 7"],f"※ 粗利益額{f0(m_gp)}千円は前期比{p1(yoy(T,'gp','m'))}・計画比{p1(vsp(T,'gp','m'))}。減益は前期に機械大口（岡野農場向け29,400千円）のあった境港（{sgn(gpdiff['境港'])}）のみで、他6部門は全て増益。境港除きは売上{p1(ex_sa_yoy)}・粗利{p1(ex_gp_yoy)}")
s=shapes_by_name(sl[3]); set_text(s["TextBox 5"],FOOT)
set_text(s["TextBox 2"],"営業所別　下期（4〜9月）累計　～売上・粗利益額～")
set_text(s["TextBox 3"],"単位: 千円　　第51期下期（4〜9月）確定　　金色列＝粗利益額（主指標）")
fill_table(sl[3],"h2")
set_text(s["TextBox 7"],f"※ 下期粗利益額{f0(h_gp)}千円・前期比{p1(yoy(T,'gp','h2'))}・計画比{p1(vsp(T,'gp','h2'))}（売上{p1(yoy(T,'sales','h2'))}・{p1(vsp(T,'sales','h2'))}）。境港は前期比{p1(yoy('境港','gp','h2'))}・計画比{p1(vsp('境港','gp','h2'))}、キョウワは前期比{p1(yoy('キョウワ','gp','h2'))}と前期を下回る")
s=shapes_by_name(sl[4]); set_text(s["TextBox 5"],FOOT)
set_text(s["TextBox 2"],"営業所別　期初来累計（第51期通期）　～売上・粗利益額～")
set_text(s["TextBox 3"],"単位: 千円　　令和7年10月〜令和8年9月（12ヶ月・第51期通期）　　金色列＝粗利益額（主指標）")
fill_table(sl[4],"ytd")
set_text(s["TextBox 7"],f"※ 第51期通期の粗利益額{f0(y_gp)}千円・前期比{p1(yoy(T,'gp','ytd'))}・計画比{p1(vsp(T,'gp','ytd'))}（売上{p1(yoy(T,'sales','ytd'))}・計画比{p1(vsp(T,'sales','ytd'))}）。境港（{p1(yoy('境港','gp','ytd'))}）・キョウワ（{p1(yoy('キョウワ','gp','ytd'))}）以外の5部門が増益")
# ---------- slide 6/7 charts
def replace_chart(slide,series):
    ch=[sh for sh in slide.shapes if sh.has_chart][0].chart
    cd=CategoryChartData(); cd.categories=OFF
    for nm,vals in series: cd.add_series(nm,vals)
    ch.replace_data(cd)
    return ch
ser6=[("単月9月",[round(yoy(o,"gp","m"),1) for o in OFF]),("下期4〜9月",[round(yoy(o,"gp","h2"),1) for o in OFF]),("通期10〜9月",[round(yoy(o,"gp","ytd"),1) for o in OFF])]
ch=replace_chart(sl[5],ser6)
mx=max(v for _,vs in ser6 for v in vs)
axmax=250.0 if mx>200 else 200.0
for el in ch._chartSpace.iter("{http://schemas.openxmlformats.org/drawingml/2006/chart}max"): el.set("val",str(axmax))
s=shapes_by_name(sl[5]); set_text(s["TextBox 5"],FOOT)
set_text(s["TextBox 7"],f"全社の粗利益額前期比は 単月{p1(yoy(T,'gp','m'))}／下期{p1(yoy(T,'gp','h2'))}／通期{p1(yoy(T,'gp','ytd'))}。単月は境港のみ、下期・通期は境港とキョウワが前期を下回る（境港通期{p1(yoy('境港','gp','ytd'))}・キョウワ通期{p1(yoy('キョウワ','gp','ytd'))}）")
ser7=[("単月9月",[round(gpdiff[o]) for o in OFF]),("下期4〜9月",[round(get(o,"gp","h2","act")-get(o,"gp","h2","prior")) for o in OFF]),("通期10〜9月",[round(get(o,"gp","ytd","act")-get(o,"gp","ytd","prior")) for o in OFF])]
ch=replace_chart(sl[6],ser7)
s=shapes_by_name(sl[6]); set_text(s["TextBox 5"],FOOT)
set_text(s["TextBox 7"],f"9月単月の粗利は{sgn(tot_gpdiff)}千円。境港{sgn(gpdiff['境港'])}（前期岡野農場向け機械大口の反動）を {joinc(plus,4)} の増益で吸収。通期では水産部{sgn(get('水産部','gp','ytd','act')-get('水産部','gp','ytd','prior'))}が最大の寄与")
# ---------- slide 8 narrative
s=shapes_by_name(sl[7]); set_text(s["TextBox 5"],FOOT)
set_text(s["TextBox 8"],f"第51期通期は粗利益額 前期比{p1(yoy(T,'gp','ytd'))}・計画比{p1(vsp(T,'gp','ytd'))}で増益着地")
set_text(s["TextBox 9"],f"・ 9月粗利益額{f0(m_gp)}千円（前期比{p1(yoy(T,'gp','m'))}・計画比{p1(vsp(T,'gp','m'))}）。売上は横ばい、増益は粗利率改善による")
set_text(s["TextBox 10"],f"・ 通期売上{f0(y_sa)}千円（前期比{p1(yoy(T,'sales','ytd'))}）・粗利{f0(y_gp)}千円（{p1(yoy(T,'gp','ytd'))}）。下期粗利は前期比{p1(yoy(T,'gp','h2'))}と上期から加速")
set_text(s["TextBox 11"],f"・ 粗利率は単月{p1(rate(T,'m'))}（前期{p1(rate(T,'m','prior'))}）・下期{p1(rate(T,'h2'))}（{p1(rate(T,'h2','prior'))}）・通期{p1(rate(T,'ytd'))}（{p1(rate(T,'ytd','prior'))}）と全期間で改善")
set_text(s["TextBox 14"],"【警戒】価格改定の検証　～数量の落ち込みはないか～")
set_text(s["TextBox 15"],f"・ 継続取扱品の単価は前年比{PVT['price_idx']-100:+.1f}%、数量は{PVT['qty_idx']:.1f}%。増益は価格改定の効果が主因")
set_text(s["TextBox 16"],f"・ 石見は単価{PVI['price_idx']-100:+.1f}%に対し数量{PVI['qty_idx']:.1f}%と減少。スチロール成形品の積水化成品西部向け等")
set_text(s["TextBox 17"],f"・ 継続取引先{PVT['cust']['n_both']}先中{PVT['cust']['down']}先は、値上げ後も売上が前年比10%超減。数量を要確認")
set_text(s["TextBox 20"],"営業所別の実態　～9月の増益はどこから来たか～")
set_text(s["TextBox 21"],f"・ 広域：粗利{p1(yoy('広域','gp','m'))}（リンガーハット東京本社{sgn(3250)}）、松江：{p1(yoy('松江','gp','m'))}（コクヨー・岡田商店等）")
set_text(s["TextBox 22"],f"・ 石見{sgn(gpdiff['石見'])}（キヌヤ）、下関{sgn(gpdiff['下関'])}（フクシン等）。境港除く6部門が増益、粗利率も全社で{rate(T,'m')-rate(T,'m','prior'):+.1f}pt")
set_text(s["TextBox 24"],f"第51期は粗利益額{f0(y_gp)}千円で増益着地。第52期は価格改定後の数量（特に石見）の確認と粗利率の維持が焦点")

import unicodedata
def w(t): return sum(1 if unicodedata.east_asian_width(c) in "FWA" else 0.5 for c in t)
limits={(1,"TextBox 8"):56,(1,"TextBox 9"):56,(1,"TextBox 7"):52,(1,"TextBox 4"):48,(1,"TextBox 13"):60,(8,"TextBox 9"):56,(8,"TextBox 10"):56,(8,"TextBox 11"):56,(8,"TextBox 15"):56,(8,"TextBox 16"):56,(8,"TextBox 17"):56,(8,"TextBox 21"):56,(8,"TextBox 22"):56,(8,"TextBox 24"):64,(8,"TextBox 14"):48,(8,"TextBox 8"):48,(8,"TextBox 14"):48,(3,"TextBox 7"):130,(4,"TextBox 7"):130,(5,"TextBox 7"):130,(6,"TextBox 7"):130,(7,"TextBox 7"):130,(2,"TextBox 3"):80}
for (si,nm),lim in limits.items():
    t=shapes_by_name(sl[si-1])[nm].text_frame.text
    flag="OK " if w(t)<=lim else "OVER"
    print(f"{flag} slide{si} {nm} width={w(t):.1f}/{lim}: {t}")
prs.save(S+"out/令和8年9月実績報告.pptx")
print("saved")
