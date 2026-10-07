# -*- coding: utf-8 -*-
# 【営業会議資料】に「通期」シート（第51期 令和7年10月〜令和8年9月）を追加
#  通期 = 下期累計シート（4〜9月）の各数値セル ＋ 上期（10〜3月：上期の営業所別ファイルを担当者コード別に集計）
import openpyxl, json, re, copy, sys
from openpyxl.styles import Font
sys.path.insert(0,'.')
import h1_parse
S="/tmp/claude-0/-home-user-claudecode/7f1b344f-da50-502a-9e34-aab355afdd38/scratchpad/"
SRC=sys.argv[1]; DST=sys.argv[2]
H1={o:h1_parse.parse(o) for o in h1_parse.FILES}
K=openpyxl.load_workbook(S+"sep/kaigi_2603.xlsx",data_only=True)["累計実績"]
KY={sec:(float(K.cell(125,c).value or 0),float(K.cell(125,c+7).value or 0)) for sec,c in (("y2",90),("y1",91),("plan",92),("act",93))}  # キョウワ(税抜)は3月営業会議資料の累計を使用
NAME2CODE={("石見","佐々木課長"):"0011",("石見","川上"):"0012",("石見","三浦"):"0013",("石見","木村課長"):"0014",("石見","その他"):"0019",("石見","自治体"):"0020",
 ("下関","高橋所長"):"0021",("下関","東野"):"0024",("下関","橋本課長"):"0025",("下関","井上課長"):"0026",("下関","中国"):"0027",("下関","自治体"):"0028",("下関","その他"):"0029",
 ("松江","飯塚主任"):"0032",("松江","前田課長代理"):"0033",("松江","原田"):"0034",("松江","自治体"):"0035",("松江","その他"):"0039",
 ("広域","馬場部長（鳥栖）"):"0042",("広域","ペーパーハグ"):"0045",("広域","角田"):"0046",("広域","馬場部長（静岡）"):"0047",("広域","馬場所長（東京）"):"0550",("広域","馬場部長（東京）"):"0552",
 ("境港","西岡"):"0903",("境港","足立次長"):"0904",("境港","藤波課長"):"0909",("境港","足立課長"):"0910",("境港","川邊主任"):"0911",("境港","白根主任"):"0914",("境港","海外"):"0918",("境港","その他"):"0919",("境港","自治体"):"0930"}
OFFICE_ROWS={"石見":(7,29),"下関":(30,55),"松江":(56,75),"広域":(76,99),"境港":(100,131)}
notes=[]
def h1(o,sh,sec,code):
    d=H1[o][sh].get(sec,{})
    if o=="境港" and code=="0919": code="0919+0930" if "0919+0930" in d else "0919"
    if o=="境港" and code=="0930" and "0919+0930" in d: return (0.0,0.0)
    return d.get(code,(0.0,0.0))
def vals(o,sh,code):
    out={}
    for col,sec in (("F","y2"),("G","y1"),("H","plan"),("I","act")): out[col]=h1(o,sh,sec,code)[0]
    for col,sec in (("M","y2"),("N","y1"),("O","plan"),("P","act")): out[col]=h1(o,sh,sec,code)[1]
    return out
def totals(o,sh):
    out={}
    for col,sec in (("F","y2"),("G","y1"),("H","plan"),("I","act")): out[col]=H1[o][sh][sec]["合計"][0]
    for col,sec in (("M","y2"),("N","y1"),("O","plan"),("P","act")): out[col]=H1[o][sh][sec]["合計"][1]
    return out
wb=openpyxl.load_workbook(SRC)
ws0=wb["下期累計"]; ws=wb.copy_worksheet(ws0); ws.title="通期"
wb._sheets.remove(ws); wb._sheets.insert(wb.sheetnames.index("下期累計")+1,ws)
ws.freeze_panes=ws0.freeze_panes; ws.sheet_view.zoomScale=ws0.sheet_view.zoomScale
for key,w in ws0.column_dimensions.items(): ws.column_dimensions[key].width=w.width
ws["A2"]="3．通期実績（第51期：令和7年10月〜令和8年9月）"
for col in ["F","M","T"]: ws[f"{col}5"]="第49期通期"
for col in ["G","N","U"]: ws[f"{col}5"]="第50期通期"
for col in ["H","I","O","P","V","W"]: ws[f"{col}5"]="第51期通期"
ws["D3"]="※上期（10〜3月）は上期の営業所別担当者別実績を担当者コード別に合算。上期は鳥栖・静岡を合算管理のため、馬場部長（鳥栖）欄に上期分を含む（静岡欄は下期分のみ）。キョウワ上期は3月営業会議資料の累計。"
ws["D3"].font=Font(size=9,color="666666")
cur=None; pend=None; changed=0
for r in range(7,144):
    d=ws.cell(r,4).value; e=ws.cell(r,5).value; b=ws.cell(r,2).value; a=ws.cell(r,1).value
    for o,(r1,r2) in OFFICE_ROWS.items():
        if r1<=r<=r2: cur=o
    def add(v):
        global changed
        for col,val in v.items():
            c=ws[f"{col}{r}"]
            if isinstance(c.value,str) and c.value.startswith("="): notes.append(f"数式セルのため加算せず {col}{r}"); continue
            c.value=(c.value or 0)+val; changed+=1
    if d and cur and (cur,d) in NAME2CODE:
        code=NAME2CODE[(cur,d)]; add(vals(cur,"全体",code)); pend=(cur,code)
    elif e=="機械" and isinstance(ws0.cell(r,6).value,(int,float)) and cur and pend:
        add(vals(pend[0],"機械",pend[1]))
    elif b=="水産部":
        add(totals("水産部","全体")); pend=None
    elif e=="荷役" and r==137:
        v1=vals("水産部","全体","0104"); v2=vals("水産部","全体","0106"); add({k:v1[k]+v2[k] for k in v1})
    elif e=="機械" and r==138:
        add(totals("水産部","機械"))
    elif a and "キョウワ" in str(a):
        add({"F":KY["y2"][0],"G":KY["y1"][0],"H":KY["plan"][0],"I":KY["act"][0],"M":KY["y2"][1],"N":KY["y1"][1],"O":KY["plan"][1],"P":KY["act"][1]})
    elif d and cur and d not in ("営業","自治体","中国事業","海外","【参考】馬場部長合計"):
        notes.append(f"未対応ラベル row {r}: {d}")
wb.save(DST); print("cells added:",changed); print(*notes,sep="\n")
