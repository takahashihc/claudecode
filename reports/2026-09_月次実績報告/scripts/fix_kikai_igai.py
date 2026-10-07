# 第51期実績: 合計ブロックの「機械以外」= 合計−機械−自治体−海外 に修正（全シート）
import openpyxl, re, sys
p=sys.argv[1]; wb=openpyxl.load_workbook(p); log=[]
for ws in wb.worksheets:
    for r in range(1,ws.max_row+1):
        if str(ws.cell(r,8).value or '').strip()!='機械以外' or str(ws.cell(r-1,4).value or '').strip()!='合計': continue
        rt,rm,rj,rk=r-1,r+1,r+2,r+3
        assert [str(ws.cell(x,8).value or '').strip() for x in (rm,rj,rk)]==['機械','自治体','海外'],(ws.title,r)
        pat=re.compile(r"=([A-Z]+)%d-(?P<c2>[A-Z]+)%d$"%(rt,rm))
        for c in range(1,ws.max_column+1):
            v=ws.cell(r,c).value
            if not isinstance(v,str): continue
            m=pat.fullmatch(v)
            if m and m.group(1)==m.group('c2'):
                col=m.group(1); new=f"={col}{rt}-{col}{rm}-{col}{rj}-{col}{rk}"; ws.cell(r,c).value=new; log.append(f"{ws.title}!{ws.cell(r,c).coordinate}: {v} -> {new}")
wb.save(p); print(len(log),'cells fixed'); print(*log,sep='\n')
