# -*- coding: utf-8 -*-
# 9月(2609) 集計: 単月 / 下期(4〜9月) / 期初来(10〜9月=第51期通期)
import openpyxl, json
S="/tmp/claude-0/-home-user-claudecode/7f1b344f-da50-502a-9e34-aab355afdd38/scratchpad/"
SEP=S+"sep/"; W=S+"work/calc/"
M=5  # Sep index (Apr=0)
OFF=["石見","下関","松江","広域","境港","水産部","キョウワ"]
files={o:f"第51期下期{o}{'営業部' if o=='広域' else ('部門別' if o=='水産部' else '営業所')}担当者別実績.xlsx" for o in ["石見","下関","松江","広域","境港"]}
files["水産部"]="第51期下期水産部部門別実績.xlsx"
codes={"石見":["0011","0012","0013","0014","0015","0016","0018","0019","0020"],
       "下関":["0021","0024","0025","0026","0027","0028","0029"],
       "松江":["0032","0033","0034","0035","0039"],
       "広域":["0042","0045","0046","0047","0550","0552"],
       "境港":["0903","0904","0909","0910","0911","0914","0918","0919","0930"],
       "水産部":["0010","0017","0031","0038","0101","0102","0104","0106"]}
raw={}
ws=openpyxl.load_workbook(SEP+"2609担当者別売上粗利益実績.xlsx",data_only=True).worksheets[0]
for row in ws.iter_rows(min_row=2,values_only=True):
    if row[0] is None: continue
    raw[str(row[0]).zfill(4)]={"name":row[1],"sales":[float(x or 0) for x in row[2:8]],"gp":[float(x or 0) for x in row[8:14]]}
allcodes=set(sum(codes.values(),[]))
print("UNMAPPED:",{c:(v["name"],round(v["sales"][M],1),round(v["gp"][M],1)) for c,v in raw.items() if c not in allcodes and c!="0502" and (any(v["sales"]) or any(v["gp"]))})
sep={o:{"sales":sum(raw[c]["sales"][M] for c in cs if c in raw),"gp":sum(raw[c]["gp"][M] for c in cs if c in raw)} for o,cs in codes.items()}
aug_raw={o:{"sales":sum(raw[c]["sales"][M-1] for c in cs if c in raw),"gp":sum(raw[c]["gp"][M-1] for c in cs if c in raw)} for o,cs in codes.items()}
def find_section(rows,prefix):
    for i,r in enumerate(rows):
        if r[2] and str(r[2]).startswith(prefix): return i
    raise Exception(prefix)
def total_row(rows,start,label="合計"):
    for r in rows[start:start+40]:
        if r[2]==label: return r
    raise Exception(label)
data={}
for o,f in files.items():
    rows=list(openpyxl.load_workbook(W+f,data_only=True)["全体"].iter_rows(values_only=True)); d={}
    for key,pre in [("prior","2．"),("plan","4．"),("act","7．")]:
        st=find_section(rows,pre); tr=total_row(rows,st)
        d[key]={"sales":[float(tr[4+i] or 0) for i in range(6)],"gp":[float(tr[11+i] or 0) for i in range(6)]}
    data[o]=d
rows=list(openpyxl.load_workbook(W+files["広域"],data_only=True)["キョウワ"].iter_rows(values_only=True))
def ky(prefix,label="合計(税抜）"):
    st=find_section(rows,prefix)
    for r in rows[st:st+8]:
        if r[2]==label: return {"sales":[float(r[4+i] or 0) for i in range(6)],"gp":[float(r[11+i] or 0) for i in range(6)]}
data["キョウワ"]={"prior":ky("2．"),"plan":ky("4．"),"act":ky("7．")}
for o in files:
    print(f"AUG check {o}: raw {aug_raw[o]['sales']:.3f}/{aug_raw[o]['gp']:.3f} vs file {data[o]['act']['sales'][M-1]:.3f}/{data[o]['act']['gp'][M-1]:.3f}")
for o in files:
    data[o]["act"]["sales"][M]=sep[o]["sales"]; data[o]["act"]["gp"][M]=sep[o]["gp"]
print("Kyowa file act Apr-Aug:",data["キョウワ"]["act"]["sales"][:5],"raw 0502:",raw["0502"]["sales"])
for i in range(6):
    data["キョウワ"]["act"]["sales"][i]=raw["0502"]["sales"][i]; data["キョウワ"]["act"]["gp"][i]=raw["0502"]["gp"][i]
# H1 (10〜3月) from 3月会議資料 累計実績
ws=openpyxl.load_workbook(SEP+"kaigi_2603.xlsx",data_only=True)["累計実績"]
h1rows={"石見":25,"下関":48,"松江":65,"広域":82,"境港":109,"水産部":117,"キョウワ":125}
h1={}
for o,r in h1rows.items():
    g=lambda c: float(ws.cell(r,c).value or 0)
    h1[o]={"sales":{"prior":g(91),"plan":g(92),"act":g(93)},"gp":{"prior":g(98),"plan":g(99),"act":g(100)}}
    print("H1",o,ws.cell(r,2).value,ws.cell(r,4).value,h1[o]["sales"]["act"],h1[o]["gp"]["act"])
# 下期4〜8月 from 営業会議資料8月 下期累計
ws=openpyxl.load_workbook(W+"【営業会議資料】令和8年8月実績.xlsx",data_only=True)["下期累計"]
lab={"石見営業所":"石見","下関営業所":"下関","松江営業所":"松江","広域営業部":"広域","境港営業所":"境港","水産部":"水産部","キョウワ（税抜）":"キョウワ"}
h2_48={}
for row in ws.iter_rows(values_only=True):
    labs=[str(v).strip() for v in row[:4] if isinstance(v,str)]
    if labs and labs[0] in lab and lab[labs[0]] not in h2_48:
        h2_48[lab[labs[0]]]={"sales":{"prior":float(row[6]),"plan":float(row[7]),"act":float(row[8])},"gp":{"prior":float(row[13]),"plan":float(row[14]),"act":float(row[15])}}
for o in OFF:
    fs=sum(data[o]["act"]["sales"][:5]); fg=sum(data[o]["act"]["gp"][:5])
    print(f"  {o}: file Apr-Aug {fs:.1f}/{fg:.1f}  kaigi {h2_48[o]['sales']['act']:.1f}/{h2_48[o]['gp']['act']:.1f}  plan file {sum(data[o]['plan']['sales'][:5]):.0f} kaigi {h2_48[o]['sales']['plan']:.0f}  prior file {sum(data[o]['prior']['sales'][:5]):.1f} kaigi {h2_48[o]['sales']['prior']:.1f}")
out={"offices":{}}
for o in OFF:
    res={}
    for m in ["sales","gp"]:
        mth={k:data[o][k][m][M] for k in ["act","plan","prior"]}
        h2={k:h2_48[o][m][k]+mth[k] for k in ["act","plan","prior"]}
        ytd={k:h2[k]+h1[o][m][k] for k in ["act","plan","prior"]}
        res[m]={"m":mth,"h2":h2,"ytd":ytd}
    out["offices"][o]=res
tot={m:{p:{k:sum(out["offices"][o][m][p][k] for o in OFF) for k in ["act","plan","prior"]} for p in ["m","h2","ytd"]} for m in ["sales","gp"]}
out["offices"]["合計"]=tot
trend={}
for i,mn in enumerate(["4月","5月","6月","7月","8月","9月"]):
    a=sum(data[o]["act"]["gp"][i] for o in OFF); p=sum(data[o]["prior"]["gp"][i] for o in OFF)
    sa=sum(data[o]["act"]["sales"][i] for o in OFF); sp=sum(data[o]["prior"]["sales"][i] for o in OFF)
    pl=sum(data[o]["plan"]["gp"][i] for o in OFF); spl=sum(data[o]["plan"]["sales"][i] for o in OFF)
    trend[mn]={"gp_yoy":a/p*100,"sales_yoy":sa/sp*100,"gp":a,"sales":sa,"gp_rate":a/sa*100,"gp_rate_prior":p/sp*100,"gp_plan":a/pl*100,"sales_plan":sa/spl*100}
out["trend"]=trend
out["raw_person"]={c:{"name":v["name"],"sales":v["sales"][M],"gp":v["gp"][M],"sales_aug":v["sales"][M-1],"gp_aug":v["gp"][M-1]} for c,v in raw.items()}
out["data"]=data; out["h1"]=h1; out["h2_48"]=h2_48
json.dump(out,open(S+"sep_final.json","w"),ensure_ascii=False,indent=1)
def pct(a,b): return a/b*100 if b else float('nan')
for p,lab_ in [("m","9月単月"),("h2","下期4〜9月"),("ytd","期初来10〜9月(通期)")]:
    print(f"\n=== {lab_} ===")
    for o in OFF+["合計"]:
        s=out["offices"][o]["sales"][p]; g=out["offices"][o]["gp"][p]
        print(f"{o:6}{s['act']:11,.0f}{pct(s['act'],s['prior']):7.1f}{pct(s['act'],s['plan']):7.1f} |{g['act']:9,.0f}{pct(g['act'],g['prior']):7.1f}{pct(g['act'],g['plan']):7.1f} | rate {pct(g['act'],s['act']):5.1f} (prior {pct(g['prior'],s['prior']):5.1f}) diff {g['act']-g['prior']:+,.0f}  prior s{s['prior']:,.0f} g{g['prior']:,.0f} plan s{s['plan']:,.0f} g{g['plan']:,.0f}")
print("\nTREND:",{k:{kk:round(vv,1) for kk,vv in v.items() if 'yoy' in kk or 'plan' in kk or 'rate' in kk} for k,v in trend.items()})
