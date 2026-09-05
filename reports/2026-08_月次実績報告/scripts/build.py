import json, copy, sys
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from lxml import etree
S="/tmp/claude-0/-home-user-claudecode/7f1b344f-da50-502a-9e34-aab355afdd38/scratchpad/"
D=json.load(open(S+"aug_final.json"))
O=D["offices"]; OFF=["石見","下関","松江","広域","境港","水産部","キョウワ"]
NS="{http://schemas.openxmlformats.org/drawingml/2006/main}"
def pct(a,b): return a/b*100
def f0(x): return f"{x:,.0f}"
def p1(x): return f"{x:.1f}%"
def sgn(x): return f"+{x:,.0f}" if x>=0 else f"▲{abs(x):,.0f}"
def get(o,m,p,k): return O[o][m][p][k]
def yoy(o,m,p): return pct(get(o,m,p,"act"),get(o,m,p,"prior"))
def vsp(o,m,p): return pct(get(o,m,p,"act"),get(o,m,p,"plan"))
def rate(o,p,k="act"): return pct(get(o,"gp",p,k),get(o,"sales",p,k))
T="合計"
# ---------- key figures
m_gp=get(T,"gp","m","act"); m_sa=get(T,"sales","m","act")
h_gp=get(T,"gp","h2","act"); h_sa=get(T,"sales","h2","act")
y_gp=get(T,"gp","ytd","act"); y_sa=get(T,"sales","ytd","act")
gpdiff={o:get(o,"gp","m","act")-get(o,"gp","m","prior") for o in OFF}
tot_gpdiff=m_gp-get(T,"gp","m","prior")
ex=D["excl_sakai_m"]
ex_sa_yoy=pct(ex["sales"]["act"],ex["sales"]["prior"]); ex_gp_yoy=pct(ex["gp"]["act"],ex["gp"]["prior"]); ex_gp_plan=pct(ex["gp"]["act"],ex["gp"]["plan"])
plus=sorted([o for o in OFF if gpdiff[o]>0],key=lambda o:-gpdiff[o])
minus=sorted([o for o in OFF if gpdiff[o]<0],key=lambda o:gpdiff[o])
def joinc(lst,n=3): return "・".join(f"{o}{sgn(gpdiff[o])}" for o in lst[:n])
print("KEY:",f0(m_gp),p1(yoy(T,'gp','m')),p1(vsp(T,'gp','m')),"| sales",f0(m_sa),p1(yoy(T,'sales','m')),p1(vsp(T,'sales','m')),"| rate",p1(rate(T,'m')),p1(rate(T,'m','prior')))
# ---------- helpers
def set_text(shape,text):
    tf=shape.text_frame
    paras=tf.paragraphs
    p=paras[0]
    runs=p.runs
    if runs:
        runs[0].text=text
        for r in runs[1:]: r._r.getparent().remove(r._r)
    else:
        p.text=text
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
def set_cell(cell,text,color=None):
    p=cell.text_frame.paragraphs[0]
    r=p.runs[0]; r.text=text
    for rr in p.runs[1:]: rr._r.getparent().remove(rr._r)
    if color:
        rPr=r._r.get_or_add_rPr()
        sf=rPr.find(NS+"solidFill")
        if sf is not None: rPr.remove(sf)
        sf=etree.SubElement(rPr,NS+"solidFill"); c=etree.SubElement(sf,NS+"srgbClr"); c.set("val",color)
        # solidFill must come before latin/ea: reorder
        for tag in ["latin","ea","cs"]:
            el=rPr.find(NS+tag)
            if el is not None: rPr.remove(el); rPr.append(el)
def ratio_color(v):
    v=round(v,1)
    return "1A6E1A" if v>=100 else ("C55A11" if v>=90 else "C0392B")
def shapes_by_name(slide): return {sh.name:sh for sh in slide.shapes}
prs=Presentation(S+"令和8年7月実績報告.pptx")
sl=prs.slides
FOOT="株式会社タカハシ包装センター | 令和8年8月実績"
# ---------- slide 1 (cover)
s=shapes_by_name(sl[0])
set_text(s["TextBox 3"],"令和8年8月　月次実績報告")
set_text(s["TextBox 4"],"第51期下期 8月実績（下期5ヶ月目）　～粗利益額を主軸とした分析～")
set_text(s["TextBox 7"],f"粗利益額: 8月 {f0(m_gp)}千円（前期比{p1(yoy(T,'gp','m'))}・計画比{p1(vsp(T,'gp','m'))}）")
set_text(s["TextBox 8"],f"・ 前期8月の境港機械大口（81,900千円）の反動で減収減益。境港除きは売上{p1(ex_sa_yoy)}・粗利{p1(ex_gp_yoy)}と増益継続")
set_text(s["TextBox 9"],f"・ 粗利は計画比{p1(vsp(T,'gp','m'))}と計画超過、粗利率{p1(rate(T,'m'))}（前期{p1(rate(T,'m','prior'))}）。増益は {joinc(plus,2)} が牽引")
set_text(s["TextBox 11"],"【警戒】石見の失速とナフサショックの先行き")
set_text(s["TextBox 12"],f"・ 石見は売上計画比{p1(vsp('石見','sales','m'))}・前期比{p1(yoy('石見','sales','m'))}と下期初の減収、粗利前期比{p1(yoy('石見','gp','m'))}は下期最低。価格改定効果の一巡を注視")
set_text(s["TextBox 13"],"分析期間: 単月（8月）／下期（4〜8月）／期初来累計（令和7年10月〜令和8年8月）")
set_text(s["TextBox 14"],"作成日: 令和8年9月5日　　資料作成者: 高橋将史")
# ---------- slide 2 (KPI cards)
s=shapes_by_name(sl[1])
set_text(s["TextBox 5"],FOOT)
cards=[("TextBox 8","TextBox 9","TextBox 11","8月 粗利益額",f0(m_gp),f"前期比 {p1(yoy(T,'gp','m'))} / 計画比 {p1(vsp(T,'gp','m'))}"),
       ("TextBox 14","TextBox 15","TextBox 17","8月 売上高",f0(m_sa),f"前期比 {p1(yoy(T,'sales','m'))} / 計画比 {p1(vsp(T,'sales','m'))}"),
       ("TextBox 20","TextBox 21","TextBox 23","8月 粗利率",f"{rate(T,'m'):.1f}",f"前期 {p1(rate(T,'m','prior'))} → 今期"),
       ("TextBox 26","TextBox 27","TextBox 29","下期 粗利益額",f0(h_gp),f"前期比 {p1(yoy(T,'gp','h2'))} / 計画比 {p1(vsp(T,'gp','h2'))}"),
       ("TextBox 32","TextBox 33","TextBox 35","下期 売上高",f0(h_sa),f"前期比 {p1(yoy(T,'sales','h2'))} / 計画比 {p1(vsp(T,'sales','h2'))}"),
       ("TextBox 38","TextBox 39","TextBox 41","下期 粗利率",f"{rate(T,'h2'):.1f}",f"前期 {p1(rate(T,'h2','prior'))} → 今期"),
       ("TextBox 44","TextBox 45","TextBox 47","期初来 粗利益額",f0(y_gp),f"前期比 {p1(yoy(T,'gp','ytd'))} / 計画比 {p1(vsp(T,'gp','ytd'))}"),
       ("TextBox 50","TextBox 51","TextBox 53","期初来 売上高",f0(y_sa),f"前期比 {p1(yoy(T,'sales','ytd'))} / 計画比 {p1(vsp(T,'sales','ytd'))}"),
       ("TextBox 56","TextBox 57","TextBox 59","期初来 粗利率",f"{rate(T,'ytd'):.1f}",f"前期 {p1(rate(T,'ytd','prior'))} → 今期")]
for a,b,c,t1,t2,t3 in cards:
    set_text(s[a],t1); set_text(s[b],t2); set_text(s[c],t3)
# ---------- slides 3-5 (tables)
names={"石見":"石見営業所","下関":"下関営業所","松江":"松江営業所","広域":"広域営業部","境港":"境港営業所","水産部":"水産部（計）","キョウワ":"キョウワ"}
def fill_table(slide,period):
    t=[sh for sh in slide.shapes if sh.has_table][0].table
    for ri,o in enumerate(OFF+[T],start=1):
        total=(o==T)
        sa=O[o]["sales"][period]; gp=O[o]["gp"][period]
        vals=[names.get(o,"合　計"),f0(sa["act"]),pct(sa["act"],sa["prior"]),pct(sa["act"],sa["plan"]),f0(gp["act"]),pct(gp["act"],gp["prior"]),pct(gp["act"],gp["plan"]),pct(gp["act"],sa["act"])]
        set_cell(t.cell(ri,0),vals[0]); set_cell(t.cell(ri,1),vals[1]); set_cell(t.cell(ri,4),vals[4])
        for ci in [2,3,5,6]:
            set_cell(t.cell(ri,ci),p1(vals[ci]),None if total else ratio_color(vals[ci]))
        set_cell(t.cell(ri,7),p1(vals[7]))
s=shapes_by_name(sl[2]); set_text(s["TextBox 5"],FOOT)
set_text(s["TextBox 2"],"営業所別　8月単月実績　～売上・粗利益額～")
fill_table(sl[2],"m")
set_text(s["TextBox 7"],f"※ 粗利益額{f0(m_gp)}千円は前期比{p1(yoy(T,'gp','m'))}だが計画比{p1(vsp(T,'gp','m'))}。減益は前期に機械大口のあった境港（{sgn(gpdiff['境港'])}）のみで、他6部門は全て増益（粗利計画比{p1(ex_gp_plan)}）。水産部は売上{p1(yoy('水産部','sales','m'))}・粗利{p1(yoy('水産部','gp','m'))}")
s=shapes_by_name(sl[3]); set_text(s["TextBox 5"],FOOT)
set_text(s["TextBox 2"],"営業所別　下期（4〜8月）累計　～売上・粗利益額～")
set_text(s["TextBox 3"],"単位: 千円　　第51期下期5ヶ月経過時点　　金色列＝粗利益額（主指標）")
fill_table(sl[3],"h2")
set_text(s["TextBox 7"],f"※ 下期粗利益額{f0(h_gp)}千円・前期比{p1(yoy(T,'gp','h2'))}（売上{p1(yoy(T,'sales','h2'))}）。境港は前期比{p1(yoy('境港','gp','h2'))}・計画比{p1(vsp('境港','gp','h2'))}、キョウワは前期比{p1(yoy('キョウワ','gp','h2'))}と前期を下回る")
s=shapes_by_name(sl[4]); set_text(s["TextBox 5"],FOOT)
set_text(s["TextBox 3"],"単位: 千円　　令和7年10月〜令和8年8月（11ヶ月）　　金色列＝粗利益額（主指標）")
fill_table(sl[4],"ytd")
set_text(s["TextBox 7"],f"※ 期初来粗利益額{f0(y_gp)}千円・前期比{p1(yoy(T,'gp','ytd'))}（売上{p1(yoy(T,'sales','ytd'))}）。粗利計画比{p1(vsp(T,'gp','ytd'))}と売上計画比{p1(vsp(T,'sales','ytd'))}を上回る")
# ---------- slide 6/7 charts
def replace_chart(slide,series):
    ch=[sh for sh in slide.shapes if sh.has_chart][0].chart
    cd=CategoryChartData(); cd.categories=OFF
    for nm,vals in series: cd.add_series(nm,vals)
    ch.replace_data(cd)
    return ch
ser6=[("単月8月",[round(yoy(o,"gp","m"),1) for o in OFF]),("下期4〜8月",[round(yoy(o,"gp","h2"),1) for o in OFF]),("期初来10〜8月",[round(yoy(o,"gp","ytd"),1) for o in OFF])]
ch=replace_chart(sl[5],ser6)
# axis max: キョウワ単月 >200 -> raise to 250
mx=max(v for _,vs in ser6 for v in vs)
axmax=250.0 if mx>200 else 200.0
for el in ch._chartSpace.iter("{http://schemas.openxmlformats.org/drawingml/2006/chart}max"): el.set("val",str(axmax))
s=shapes_by_name(sl[5]); set_text(s["TextBox 5"],FOOT)
set_text(s["TextBox 7"],f"全社の粗利益額前期比は 単月{p1(yoy(T,'gp','m'))}／下期{p1(yoy(T,'gp','h2'))}／期初来{p1(yoy(T,'gp','ytd'))}。単月は境港のみ、下期は境港・キョウワ、期初来はキョウワのみ前期を下回る（境港期初来は{p1(yoy('境港','gp','ytd'))}）")
ser7=[("単月8月",[round(gpdiff[o]) for o in OFF]),("下期4〜8月",[round(get(o,"gp","h2","act")-get(o,"gp","h2","prior")) for o in OFF]),("期初来10〜8月",[round(get(o,"gp","ytd","act")-get(o,"gp","ytd","prior")) for o in OFF])]
ch=replace_chart(sl[6],ser7)
s=shapes_by_name(sl[6]); set_text(s["TextBox 5"],FOOT)
set_text(s["TextBox 7"],f"8月単月の粗利は{sgn(tot_gpdiff)}千円。境港{sgn(gpdiff['境港'])}（前期ニッスイサーモン向け機械大口の反動）を {joinc(plus,4)} の増益で大半を吸収")
# ---------- slide 8 narrative
s=shapes_by_name(sl[7]); set_text(s["TextBox 5"],FOOT)
tr=D["trend"]
set_text(s["TextBox 8"],"粗利益額は前期比減も計画超過　～境港の前期特殊要因を除けば増益基調は継続～")
set_text(s["TextBox 9"],f"・ 8月粗利益額 {f0(m_gp)}千円（前期比{p1(yoy(T,'gp','m'))}・計画比{p1(vsp(T,'gp','m'))}）。売上は前期比{p1(yoy(T,'sales','m'))}・計画比{p1(vsp(T,'sales','m'))}")
set_text(s["TextBox 10"],f"・ 前期8月の境港はニッスイサーモン向け機械81,900千円を含む。境港除き6部門は売上{p1(ex_sa_yoy)}・粗利{p1(ex_gp_yoy)}・粗利計画比{p1(ex_gp_plan)}")
set_text(s["TextBox 11"],f"・ 粗利率は単月{p1(rate(T,'m'))}（前期{p1(rate(T,'m','prior'))}）・下期{p1(rate(T,'h2'))}（{p1(rate(T,'h2','prior'))}）・期初来{p1(rate(T,'ytd'))}（{p1(rate(T,'ytd','prior'))}）と全期間で改善")
set_text(s["TextBox 14"],"【警戒】石見の失速とナフサショック　～価格改定効果の一巡が数字に出始めた～")
set_text(s["TextBox 15"],f"・ 石見は売上計画比{p1(vsp('石見','sales','m'))}・前期比{p1(yoy('石見','sales','m'))}で下期初の減収。粗利前期比は4〜7月の2桁増から{p1(yoy('石見','gp','m'))}へ鈍化")
set_text(s["TextBox 16"],f"・ 全社粗利前期比は6月{tr['6月']['gp_yoy']:.1f}%→7月{tr['7月']['gp_yoy']:.1f}%→8月{yoy(T,'gp','m'):.1f}%（境港除き{ex_gp_yoy:.1f}%）。売上計画比{p1(vsp(T,'sales','m'))}は下期初の計画割れ")
set_text(s["TextBox 17"],"・ 原料再騰なら転嫁が追いつかず粗利率は急落しうる。価格改定効果剥落後の「実力の粗利益額」確保を最優先に")
set_text(s["TextBox 20"],"営業所別の実態　～売上と粗利益額の乖離～")
set_text(s["TextBox 21"],f"・ 広域：粗利前期比{p1(yoy('広域','gp','m'))}（リンガーハット向け牽引）、松江：{p1(yoy('松江','gp','m'))}（岡田商店・JAしまね等）と粗利率改善を伴う増益")
set_text(s["TextBox 22"],f"・ キョウワ：売上{f0(get('キョウワ','sales','m','act'))}千円（前期比{p1(yoy('キョウワ','sales','m'))}）と大幅増だが粗利率{p1(rate('キョウワ','m'))}（前期{p1(rate('キョウワ','m','prior'))}）。下期粗利前期比{p1(yoy('キョウワ','gp','h2'))}は前期割れ")
set_text(s["TextBox 24"],"最重要指標は粗利益額。境港の特殊要因を除けば増益だが、石見の失速と伸び率鈍化は価格改定効果一巡の兆候。粗利率の変調を月次で監視")

import unicodedata
def w(t): return sum(1 if unicodedata.east_asian_width(c) in "FWA" else 0.5 for c in t)
limits={(1,"TextBox 8"):56,(1,"TextBox 9"):56,(1,"TextBox 12"):56,(1,"TextBox 7"):45,(8,"TextBox 9"):62,(8,"TextBox 10"):62,(8,"TextBox 11"):62,(8,"TextBox 15"):62,(8,"TextBox 16"):62,(8,"TextBox 17"):62,(8,"TextBox 21"):62,(8,"TextBox 22"):62,(8,"TextBox 24"):64,(8,"TextBox 8"):48,(8,"TextBox 14"):48,(3,"TextBox 7"):130,(4,"TextBox 7"):130,(5,"TextBox 7"):130,(6,"TextBox 7"):130,(7,"TextBox 7"):130}
for (si,nm),lim in limits.items():
    t=shapes_by_name(sl[si-1])[nm].text_frame.text
    flag="OK " if w(t)<=lim else "OVER"
    print(f"{flag} slide{si} {nm} width={w(t):.1f}/{lim}: {t}")
prs.save(S+"令和8年8月実績報告.pptx")
print("saved")
