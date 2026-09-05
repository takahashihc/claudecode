# -*- coding: utf-8 -*-
import openpyxl, copy
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
S="/tmp/claude-0/-home-user-claudecode/7f1b344f-da50-502a-9e34-aab355afdd38/scratchpad/"; OUT=S+"out/"
# ---------------- A. 仕入先集計 (template: 2607仕入先実績_集計.xlsx)
tpl=openpyxl.load_workbook(S+"2607仕入先実績_集計.xlsx"); tws=tpl.active
sty_data=[copy.copy(tws.cell(4,c)._style) for c in range(1,12)]
sty_total=[copy.copy(tws.cell(tws.max_row,c)._style) for c in range(1,12)]
raw=openpyxl.load_workbook(S+"2608仕入先実績.xlsx",data_only=True).active
agg={}
for r in range(2,raw.max_row+1):
    code=str(raw.cell(r,1).value or "").strip()
    if not code: continue
    key=code[:11]
    a=agg.setdefault(key,{"codes":[],"name":raw.cell(r,2).value,"v":[0]*6})
    a["codes"].append(code)
    for i in range(6): a["v"][i]+=float(raw.cell(r,3+i).value or 0)
rows=sorted(agg.values(),key=lambda a:(-(a["v"][3]+a["v"][4]+a["v"][5]),a["codes"][0]))
wb=openpyxl.load_workbook(S+"2607仕入先実績_集計.xlsx"); ws=wb.active
ws.delete_rows(4,ws.max_row-3)
ws["A1"]="2026年8月　仕入先別実績（債務部門統合・累計純仕入額順）"
r=4
for a in rows:
    v=a["v"]; vals=["/".join(a["codes"]),a["name"],v[0],v[1],v[2],v[0]+v[1]+v[2],v[3],v[4],v[5],v[3]+v[4]+v[5],len(a["codes"])]
    for c,val in enumerate(vals,start=1):
        cell=ws.cell(r,c); cell.value=val; cell._style=copy.copy(sty_data[c-1])
    r+=1
tot=[sum(a["v"][i] for a in rows) for i in range(6)]
vals=[None,"合計",tot[0],tot[1],tot[2],tot[0]+tot[1]+tot[2],tot[3],tot[4],tot[5],tot[3]+tot[4]+tot[5],None]
for c,val in enumerate(vals,start=1):
    cell=ws.cell(r,c); cell.value=val; cell._style=copy.copy(sty_total[c-1])
wb.save(OUT+"2608仕入先実績_集計.xlsx"); print("saved 仕入先集計 rows",len(rows),"total 当月仕入",tot[0])
# ---------------- B. 担当者別得意先別 2_new (build inside July template workbook to keep styles)
wb=openpyxl.load_workbook(S+"2607担当者別得意先別売上粗利益実績2_new.xlsx"); ws=wb.active
sty_group=[copy.copy(ws.cell(4,c)._style) for c in range(1,31)]
sty_row=[copy.copy(ws.cell(5,c)._style) for c in range(1,31)]
sty_sub=[copy.copy(ws.cell(9,c)._style) for c in range(1,31)]
ws.delete_rows(4,ws.max_row-3)
ws["A1"]="2026年8月　担当者別得意先別売上粗利益"
raw=openpyxl.load_workbook(S+"2608担当者別得意先別売上粗利益実績.xlsx",data_only=True).active
recs=[]
for r in range(2,raw.max_row+1):
    v=[raw.cell(r,c).value for c in range(1,23)]
    if v[0] is None: continue
    recs.append(v)
groups={}
for v in recs: groups.setdefault(str(v[2]).zfill(4),{"name":v[3],"rows":[]})["rows"].append(v)
def nz(x): return None if (x is None or x==0) else x
def rate(g,s): return None if not s else round(g/s*100,1)
def ratio(a,b): return None if (not a or not b) else round(a/b*100,1)
def block(s,g,ps,pg,pps,ppg,r_raw=None,pr_raw=None,ppr_raw=None,sub=False):
    f=(lambda x:x) if sub else nz
    return [f(s),f(g),(r_raw if (r_raw is not None and not sub) else rate(g,s)) if s else None,
            f(ps),f(pg),(pr_raw if (pr_raw is not None and not sub) else rate(pg,ps)) if ps else None,
            ratio(s,ps),ratio(g,pg),
            f(pps),f(ppg),(ppr_raw if (ppr_raw is not None and not sub) else rate(ppg,pps)) if pps else None,
            ratio(s,pps),ratio(g,ppg)]
r=4
for code in sorted(groups):
    grp=groups[code]
    ws.cell(r,1).value=f"{code}  {grp['name']}"
    for c in range(1,31): ws.cell(r,c)._style=copy.copy(sty_group[c-1])
    r+=1
    tot=[0]*12
    for v in sorted(grp["rows"],key=lambda v:-(v[19] or 0)):
        E,F,G,H,I,J,K,Lg,M,N,O,P,Q,R,Sx,T,U,V=[(x or 0) for x in v[4:22]]
        m=block(K,Lg,H,I,E,F,M,J,G); cum=block(T,U,Q,R,N,O,V,Sx,P)
        vals=[code,grp["name"],str(v[0]),v[1]]+m+cum
        for c,val in enumerate(vals,start=1):
            cell=ws.cell(r,c); cell.value=val; cell._style=copy.copy(sty_row[c-1])
        for i,x in enumerate([K,Lg,H,I,E,F,T,U,Q,R,N,O]): tot[i]+=x
        r+=1
    m=block(tot[0],tot[1],tot[2],tot[3],tot[4],tot[5],sub=True); cum=block(tot[6],tot[7],tot[8],tot[9],tot[10],tot[11],sub=True)
    vals=[None,"小計",None,None]+m+cum
    for c,val in enumerate(vals,start=1):
        cell=ws.cell(r,c); cell.value=val; cell._style=copy.copy(sty_sub[c-1])
    r+=1
ws.freeze_panes="E4"
wb.save(OUT+"2608担当者別得意先別売上粗利益実績2_new.xlsx"); print("saved 2_new rows",r-1,"groups",len(groups))
