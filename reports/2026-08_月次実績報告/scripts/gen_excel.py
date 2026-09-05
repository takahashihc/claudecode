# -*- coding: utf-8 -*-
import openpyxl, re, copy, json, os
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter as L
S="/tmp/claude-0/-home-user-claudecode/7f1b344f-da50-502a-9e34-aab355afdd38/scratchpad/"
OUT=S+"out/"
AUG=4  # month index Apr=0..Sep=5
files={"石見":"iwami_shimoki.xlsx","下関":"第51期下期下関営業所担当者別実績.xlsx","松江":"第51期下期松江営業所担当者別実績.xlsx","広域":"第51期下期広域営業部担当者別実績.xlsx","境港":"第51期下期境港営業所担当者別実績.xlsx","水産部":"第51期下期水産部部門別実績.xlsx"}
outnames={"石見":"第51期下期石見営業所担当者別実績.xlsx","下関":"第51期下期下関営業所担当者別実績.xlsx","松江":"第51期下期松江営業所担当者別実績.xlsx","広域":"第51期下期広域営業部担当者別実績.xlsx","境港":"第51期下期境港営業所担当者別実績.xlsx","水産部":"第51期下期水産部部門別実績.xlsx"}
flags=[]
# ---------------- raw Aug data
def load_raw(f):
    ws=openpyxl.load_workbook(S+f,data_only=True).worksheets[0]; d={}
    for row in ws.iter_rows(min_row=2,values_only=True):
        if row[0] is None: continue
        d[str(row[0]).zfill(4)]={"sales":[float(x or 0) for x in row[2:7]],"gp":[float(x or 0) for x in row[7:12]]}
    return d
RAW=load_raw("担当者別売上粗利益実績202608.xlsx"); RAWM=load_raw("担当者別売上粗利益実績機械202608.xlsx")
def raw(code,kind,m=AUG,src=None):
    src=src or RAW
    return src.get(code,{"sales":[0]*5,"gp":[0]*5})[kind][m]
# Ringer Hut Aug from 得意先別
wc=openpyxl.load_workbook(S+"2608担当者別得意先別売上粗利益実績.xlsx",data_only=True).active
ringer={"sales":0.0,"gp":0.0}
for r in range(2,wc.max_row+1):
    n=str(wc.cell(r,2).value or "")
    if "ﾘﾝｶﾞｰ" in n or "リンガー" in n:
        ringer["sales"]+=(wc.cell(r,11).value or 0)/1000; ringer["gp"]+=(wc.cell(r,12).value or 0)/1000
print("Ringer Aug:",ringer)
# ---------------- office data model (from formulas-evaluated cached values)
SEC_KEYS={"1":"y2024","2":"y2025","4":"plan","7":"act"}
def sec_start(ws):
    secs={}
    for r in range(1,ws.max_row+1):
        c=ws.cell(r,3).value
        if isinstance(c,str) and len(c)>=2 and c[1] in "．,":
            k=c[0].translate(str.maketrans("１２３４５６７８９","123456789"))
            secs[k]=r
    return secs
def parse_office(o):
    """returns {sheet:{key:{label:{'sales':[6],'gp':[6]}}}} using cached values"""
    wb=openpyxl.load_workbook(S+files[o],data_only=True); res={}
    for sh in ["全体","機械"]:
        ws=wb[sh]; secs=sec_start(ws); res[sh]={}
        order=sorted(secs.items(),key=lambda x:x[1])
        for i,(k,st) in enumerate(order):
            if k not in SEC_KEYS: continue
            en=order[i+1][1] if i+1<len(order) else ws.max_row+1
            d={}
            for r in range(st+1,en):
                c=ws.cell(r,3).value
                if c is None: continue
                lab=str(c).strip()
                m=re.match(r"(\d{4})",lab)
                key=m.group(1) if m else lab
                d[key]={"sales":[float(ws.cell(r,5+j).value or 0) if not isinstance(ws.cell(r,5+j).value,str) else 0 for j in range(6)],
                        "gp":[float(ws.cell(r,12+j).value or 0) if not isinstance(ws.cell(r,12+j).value,str) else 0 for j in range(6)],"row":r}
            res[sh][SEC_KEYS[k]]=d
    return res
OFFICE={o:parse_office(o) for o in files}
# sanity
print("石見 全体 plan 合計 Aug:",OFFICE["石見"]["全体"]["plan"]["合計"]["sales"][AUG], " 2025 合計:",OFFICE["石見"]["全体"]["y2025"]["合計"]["sales"][AUG])
# Kyowa sheet model
wk=openpyxl.load_workbook(S+files["広域"],data_only=True)["キョウワ"]
def kyowa_rows():
    secs=sec_start(wk); order=sorted(secs.items(),key=lambda x:x[1]); res={}
    for i,(k,st) in enumerate(order):
        en=order[i+1][1] if i+1<len(order) else wk.max_row+1
        for r in range(st+1,en):
            c=wk.cell(r,3).value
            if c and "税抜" in str(c):
                res[k]={"sales":[float(wk.cell(r,5+j).value or 0) for j in range(6)],"gp":[float(wk.cell(r,12+j).value or 0) for j in range(6)],"row":r}
    return res
KY=kyowa_rows(); print("Kyowa secs:",{k:(v['row'],v['sales'][AUG]) for k,v in KY.items()})
def kyowa(key,kind,m=AUG):
    k={"y2024":"1","y2025":"2","plan":"4","act":"7"}[key]
    if key=="act":
        return round(raw("0502",kind,m))
    return KY[k][kind][m]
# actual by code (Apr-Jul from office file act section, Aug from raw)
def act(o,sh,code,kind,m):
    if m==AUG:
        src=RAWM if sh=="機械" else RAW
        return raw(code,kind,m,src)
    return OFFICE[o][sh]["act"].get(code,{"sales":[0]*6,"gp":[0]*6})[kind][m]
def hist(o,sh,key,code,kind,m):
    return OFFICE[o][sh][key].get(code,{"sales":[0]*6,"gp":[0]*6})[kind][m]
def codes_in(o,sh="全体"):
    return [k for k in OFFICE[o][sh]["act"] if re.match(r"\d{4}",k)]
def office_total(o,sh,key,kind,m):
    if key=="act":
        if m==AUG: return sum(act(o,sh,c,kind,m) for c in codes_in(o,sh))
        return OFFICE[o][sh]["act"]["合計"][kind][m]
    return OFFICE[o][sh][key]["合計"][kind][m]
# ======================= 1. update office files
def translate_fill(ws,r,cols=(9,16,23)):
    """for target col, if empty and left neighbour has formula -> translate"""
    for c in cols:
        cell=ws.cell(r,c); left=ws.cell(r,c-1)
        if cell.value is None and isinstance(left.value,str) and left.value.startswith("="):
            cell.value=Translator(left.value,origin=left.coordinate).translate_formula(cell.coordinate)
            cell.number_format=left.number_format; cell._style=copy.copy(left._style)
def update_office(o):
    wb=openpyxl.load_workbook(S+files[o])
    for sh in ["全体","機械","包装資材"]:
        ws=wb[sh]; secs=sec_start(ws); order=sorted(secs.items(),key=lambda x:x[1])
        for i,(k,st) in enumerate(order):
            en=order[i+1][1] if i+1<len(order) else ws.max_row+1
            if k not in ("7","8","9","5","6"): continue
            for r in range(st+1,en):
                c=ws.cell(r,3).value
                if c is None: continue
                lab=str(c).strip(); m=re.match(r"(\d{4})",lab)
                if k=="7" and sh in ("全体","機械") and m and not isinstance(ws.cell(r,8).value,str):
                    code=m.group(1); src=RAWM if sh=="機械" else RAW
                    ws.cell(r,9).value=raw(code,"sales",AUG,src); ws.cell(r,16).value=raw(code,"gp",AUG,src)
                    for cc in (9,16):
                        ws.cell(r,cc).number_format=ws.cell(r,cc-1).number_format
                if k=="7" and lab=="内リンガーハット":
                    if sh=="全体":
                        ws.cell(r,9).value=round(ringer["sales"],3); ws.cell(r,16).value=round(ringer["gp"],3)
                    elif sh=="機械":
                        flags.append(f"{o} 機械!{ws.cell(r,9).coordinate}/{ws.cell(r,16).coordinate} 内リンガーハット機械 8月: 機械分得意先別データ未提供のため空欄")
                if k=="7" and lab.startswith("【参考】大海分"):
                    flags.append(f"{o} {sh}!{ws.cell(r,9).coordinate} 【参考】大海分 8月: 算出元不明のため空欄")
                    continue
                translate_fill(ws,r)
    if o=="広域":
        for sh in ["全体","機械","包装資材"]:
            ws=wb[sh]
            for r in range(1,ws.max_row+1):
                c=str(ws.cell(r,3).value or "")
                if c.startswith("0552"): ws.cell(r,4).value="馬場部長（東京）"
                elif c.startswith("0550"): ws.cell(r,4).value="馬場所長（東京）"
        ws=wb["キョウワ"]; secs=sec_start(ws); order=sorted(secs.items(),key=lambda x:x[1])
        for i,(k,st) in enumerate(order):
            en=order[i+1][1] if i+1<len(order) else ws.max_row+1
            for r in range(st+1,en):
                c=ws.cell(r,3).value
                if c and "税抜" in str(c):
                    if k=="7":
                        ws.cell(r,9).value=round(raw("0502","sales")); ws.cell(r,16).value=round(raw("0502","gp"))
                    translate_fill(ws,r)
    wb.save(OUT+outnames[o]); print("saved office",o)
for o in files: update_office(o)
# ======================= 2. 営業会議資料 8月
NAME2CODE={("石見","佐々木課長"):"0011",("石見","川上"):"0012",("石見","三浦"):"0013",("石見","木村課長"):"0014",("石見","その他"):"0019",("石見","自治体"):"0020",
 ("下関","高橋所長"):"0021",("下関","東野"):"0024",("下関","橋本課長"):"0025",("下関","井上課長"):"0026",("下関","中国"):"0027",("下関","自治体"):"0028",("下関","その他"):"0029",
 ("松江","飯塚主任"):"0032",("松江","前田課長代理"):"0033",("松江","原田"):"0034",("松江","自治体"):"0035",("松江","その他"):"0039",
 ("広域","馬場部長（鳥栖）"):"0042",("広域","ペーパーハグ"):"0045",("広域","角田"):"0046",("広域","馬場部長（静岡）"):"0047",("広域","馬場所長（東京）"):"0550",("広域","馬場部長（東京）"):"0552",
 ("境港","西岡"):"0903",("境港","足立次長"):"0904",("境港","藤波課長"):"0909",("境港","足立課長"):"0910",("境港","川邊主任"):"0911",("境港","白根主任"):"0914",("境港","海外"):"0918",("境港","その他"):"0919",("境港","自治体"):"0930"}
OFFICE_ROWS={"石見":(7,29),"下関":(30,55),"松江":(56,75),"広域":(76,99),"境港":(100,131)}
def kaigi_values(o,code,sh,months):
    """returns dict F,G,H,I,M,N,O,P summing months"""
    v={}
    for col,key in [("F","y2024"),("G","y2025"),("H","plan")]:
        v[col]=sum(hist(o,sh,key,code,"sales",m) for m in months)
    v["I"]=sum(act(o,sh,code,"sales",m) for m in months)
    for col,key in [("M","y2024"),("N","y2025"),("O","plan")]:
        v[col]=sum(hist(o,sh,key,code,"gp",m) for m in months)
    v["P"]=sum(act(o,sh,code,"gp",m) for m in months)
    return v
def total_values(o,sh,months):
    v={}
    for col,key in [("F","y2024"),("G","y2025"),("H","plan")]: v[col]=sum(office_total(o,sh,key,"sales",m) for m in months)
    v["I"]=sum(office_total(o,sh,"act","sales",m) for m in months)
    for col,key in [("M","y2024"),("N","y2025"),("O","plan")]: v[col]=sum(office_total(o,sh,key,"gp",m) for m in months)
    v["P"]=sum(office_total(o,sh,"act","gp",m) for m in months)
    return v
def fill_kaigi_sheet(ws,months,label_m,label_h):
    # headers
    for col in ["F","M","T"]: ws[f"{col}5"]=f"令和6年{label_m}"
    for col in ["G","N","U"]: ws[f"{col}5"]=f"令和7年{label_m}"
    for col in ["H","I","O","P","V","W"]: ws[f"{col}5"]=f"令和8年{label_m}"
    ws["A2"]=label_h
    cur=None
    for r in range(7,144):
        d=ws.cell(r,4).value; e=ws.cell(r,5).value; b=ws.cell(r,2).value; a=ws.cell(r,1).value
        for o,(r1,r2) in OFFICE_ROWS.items():
            if r1<=r<=r2: cur=o
        def put(v):
            for col,val in v.items(): ws[f"{col}{r}"]=val
        if d and cur and (cur,d) in NAME2CODE:
            code=NAME2CODE[(cur,d)]; put(kaigi_values(cur,code,"全体",months)); pend=(cur,code)
        elif e=="機械" and isinstance(ws.cell(r,6).value,(int,float)) and cur and 'pend' in dir() and pend:
            put(kaigi_values(pend[0],pend[1],"機械",months))
        elif b=="水産部":
            put(total_values("水産部","全体",months)); pend=None
        elif e=="荷役" and r in (137,):
            v1=kaigi_values("水産部","0104","全体",months); v2=kaigi_values("水産部","0106","全体",months); put({k:v1[k]+v2[k] for k in v1})
        elif e=="機械" and r in (138,):
            put(total_values("水産部","機械",months))
        elif a and "キョウワ" in str(a):
            put({"F":sum(kyowa("y2024","sales",m) for m in months),"G":sum(kyowa("y2025","sales",m) for m in months),"H":sum(kyowa("plan","sales",m) for m in months),"I":sum(kyowa("act","sales",m) for m in months),
                 "M":sum(kyowa("y2024","gp",m) for m in months),"N":sum(kyowa("y2025","gp",m) for m in months),"O":sum(kyowa("plan","gp",m) for m in months),"P":sum(kyowa("act","gp",m) for m in months)})
        elif d and cur and d not in ("営業","自治体","中国事業","海外","【参考】馬場部長合計"):
            flags.append(f"営業会議資料 row {r}: 未対応ラベル {d}")
def fix_baba(ws):
    import re as _re
    ws["D91"]="馬場部長（東京）"; ws["D88"]="馬場所長（東京）"
    for c in range(6,24):
        for r,add in [(94,91),(96,93)]:
            v=ws.cell(r,c).value
            if isinstance(v,str) and _re.fullmatch(r"=([A-Z]+)\d+\+\1\d+\+\1\d+",v):
                col=_re.match(r"=([A-Z]+)",v).group(1); ws.cell(r,c).value=v+f"+{col}{add}"
    ws["J94"]="=I94-H94"
def make_kaigi():
    wb=openpyxl.load_workbook(S+"【営業会議資料】令和8年7月実績.xlsx")
    fix_baba(wb["単月"]); fix_baba(wb["下期累計"])
    fill_kaigi_sheet(wb["単月"],[AUG],"8月","1．単月実績")
    fill_kaigi_sheet(wb["下期累計"],list(range(0,AUG+1)),"4〜8月","2．下期累計実績（4〜8月）")
    wb.save(OUT+"【営業会議資料】令和8年8月実績.xlsx"); print("saved kaigi")
make_kaigi()
# ======================= 3. 第51期8月実績
def make_jisseki():
    wb=openpyxl.load_workbook(S+"51ki_07_jisseki.xlsx")
    ws7=wb["7月"]
    bd=wb.copy_worksheet(ws7); bd.title="BACKDATA7月"
    # move BACKDATA7月 right after 期初来累計
    wb._sheets.remove(bd); wb._sheets.insert(wb.sheetnames.index("期初来累計")+1,bd)
    ws7.title="8月"; ws=ws7
    ws["I3"]="2024年8月売上実績"; ws["J3"]="2025年8月売上実績"; ws["K3"]="2026年8月売上計画"; ws["L3"]="2026年8月売上実績"
    ws["P3"]="2024年8月粗利益実績"; ws["Q3"]="2025年8月粗利益実績"; ws["R3"]="2026年8月粗利益計画"; ws["S3"]="2026年8月粗利益実績"
    ws["W3"]="2024年8月粗利益率実績"; ws["X3"]="2025年8月粗利益率実績"; ws["Y3"]="2026年8月粗利益率計画"; ws["Z3"]="2026年8月粗利益率実績"
    def put(r,vals):  # vals: (s2024,s2025,splan,sact,g2024,g2025,gplan,gact)
        for col,v in zip(["I","J","K","L","P","Q","R","S"],vals): ws[f"{col}{r}"]=v
    def tot(o,sh): return tuple(office_total(o,sh,k,"sales",AUG) for k in ["y2024","y2025","plan","act"])+tuple(office_total(o,sh,k,"gp",AUG) for k in ["y2024","y2025","plan","act"])
    def per(o,code,sh="全体"): return tuple(hist(o,sh,k,code,"sales",AUG) for k in ["y2024","y2025","plan"])+(act(o,sh,code,"sales",AUG),)+tuple(hist(o,sh,k,code,"gp",AUG) for k in ["y2024","y2025","plan"])+(act(o,sh,code,"gp",AUG),)
    put(5,tot("石見","全体")); put(8,tot("石見","機械")); put(9,per("石見","0020"))
    put(10,tot("下関","全体")); put(13,tot("下関","機械")); put(14,per("下関","0028")); put(15,per("下関","0027"))
    put(16,tot("松江","全体")); put(19,tot("松江","機械")); put(20,per("松江","0035"))
    put(21,tot("広域","全体")); put(23,tot("広域","機械"))
    rg=OFFICE["広域"]["全体"]; rgm=OFFICE["広域"]["機械"]
    def rl(d,key):
        for k in d[key]:
            if "リンガー" in k and "以外" not in k: return d[key][k]
    put(24,(rl(rg,"y2024")["sales"][AUG],rl(rg,"y2025")["sales"][AUG],rl(rg,"plan")["sales"][AUG],round(ringer["sales"],3),rl(rg,"y2024")["gp"][AUG],rl(rg,"y2025")["gp"][AUG],rl(rg,"plan")["gp"][AUG],round(ringer["gp"],3)))
    put(26,(rl(rgm,"y2024")["sales"][AUG],rl(rgm,"y2025")["sales"][AUG],rl(rgm,"plan")["sales"][AUG],None,rl(rgm,"y2024")["gp"][AUG],rl(rgm,"y2025")["gp"][AUG],rl(rgm,"plan")["gp"][AUG],None))
    flags.append("第51期8月実績 8月!L26/S26 リンガーハット機械 8月実績: 機械分得意先別データ未提供のため空欄")
    put(30,tot("境港","全体")); put(33,tot("境港","機械")); put(34,per("境港","0930")); put(35,per("境港","0918"))
    put(41,tot("水産部","全体")); 
    v1=per("水産部","0104"); v2=per("水産部","0106"); put(43,tuple(a+b for a,b in zip(v1,v2)))
    put(44,tot("水産部","機械"))
    put(50,(kyowa("y2024","sales"),kyowa("y2025","sales"),kyowa("plan","sales"),kyowa("act","sales"),kyowa("y2024","gp"),kyowa("y2025","gp"),kyowa("plan","gp"),kyowa("act","gp")))
    # 下期 formulas: '7月'!X -> '8月'!X+BACKDATA7月!X
    wsd=wb["下期"]
    for row in wsd.iter_rows():
        for c in row:
            if isinstance(c.value,str) and "'7月'!" in c.value:
                c.value=re.sub(r"'7月'!(\$?[A-Z]+\$?\d+)",r"'8月'!\1+BACKDATA7月!\1",c.value)
    # any other sheet referencing '7月'
    for w in wb.worksheets:
        if w.title in ("下期",): continue
        for row in w.iter_rows():
            for c in row:
                if isinstance(c.value,str) and "'7月'!" in c.value: flags.append(f"第51期8月実績 {w.title}!{c.coordinate} still references '7月'")
    wb.save(OUT+"第51期8月実績.xlsx"); print("saved jisseki")
make_jisseki()
json.dump(flags,open(OUT+"flags.json","w"),ensure_ascii=False,indent=1)
print("FLAGS:",*flags,sep="\n ")
