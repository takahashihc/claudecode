import openpyxl, re, copy
from openpyxl.utils import get_column_letter as L
OUT="/tmp/claude-0/-home-user-claudecode/7f1b344f-da50-502a-9e34-aab355afdd38/scratchpad/out/"
names=["第51期下期石見営業所担当者別実績.xlsx","第51期下期下関営業所担当者別実績.xlsx","第51期下期松江営業所担当者別実績.xlsx","第51期下期広域営業部担当者別実績.xlsx","第51期下期境港営業所担当者別実績.xlsx","第51期下期水産部部門別実績.xlsx"]
def sec_start(ws):
    secs={}
    for r in range(1,ws.max_row+1):
        c=ws.cell(r,3).value
        if isinstance(c,str) and len(c)>=2 and c[1] in "．,":
            secs[c[0].translate(str.maketrans("１２３４５６７８９","123456789"))]=r
    return secs
log=[]
for n in names:
    wb=openpyxl.load_workbook(OUT+n)
    for sh in ["全体","機械","包装資材"]:
        ws=wb[sh]; secs=sec_start(ws)
        if not all(k in secs for k in "1247"): continue
        order=sorted(secs.items(),key=lambda x:x[1])
        ends={k:(order[i+1][1] if i+1<len(order) else ws.max_row+1) for i,(k,st) in enumerate(order)}
        def has_data(r):  # row has month cells
            return any(ws.cell(r,c).value is not None for c in range(5,11))
        for k in ["7","8","9"]:
            if k not in secs: continue
            for r in range(secs[k]+1,ends[k]):
                c3=ws.cell(r,3).value
                if c3 is None or not has_data(r): continue
                lab=str(c3)
                if lab.startswith("【参考】"): continue
                # K / R sums
                for col,rng in [(11,"E{r}:J{r}"),(18,"L{r}:Q{r}")]:
                    cell=ws.cell(r,col)
                    if cell.value is None:
                        cell.value="=SUM("+rng.format(r=r)+")"
                        cell._style=copy.copy(ws.cell(r,col-1)._style)
                        log.append(f"{n[6:9]} {sh}!{cell.coordinate} 追加 {cell.value}  ({lab} {ws.cell(r,4).value or ''})")
                # rate columns S..Y in sections 8/9 for person rows (formula pattern check)
                if k in ("8","9") and re.match(r"\d{4}",lab):
                    base=r-secs[k]; r7=secs["7"]+base; rb=(secs["4"] if k=="8" else secs["2"])+base
                    for col in range(19,26):
                        exp=f"={L(col)}{r7}-{L(col)}{rb}"
                        cell=ws.cell(r,col); cur=cell.value
                        if cur is None: continue
                        if cur!=exp and re.fullmatch(r"=[A-Z]+\d+-[A-Z]+\d+",str(cur)):
                            log.append(f"{n[6:9]} {sh}!{cell.coordinate} 修正 {cur} -> {exp}  ({lab} {ws.cell(r,4).value or ''})")
                            cell.value=exp
    wb.save(OUT+n)
print(len(log),"changes"); print(*log,sep="\n")
