# -*- coding: utf-8 -*-
import xlrd, json, openpyxl
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference
S="/tmp/claude-0/-home-user-claudecode/7f1b344f-da50-502a-9e34-aab355afdd38/scratchpad/"; OUT=S+"out/"
D=json.load(open(S+"aug_final.json"))["offices"]
def load_raw(f):
    ws=openpyxl.load_workbook(S+f,data_only=True).worksheets[0]; d={}
    for row in ws.iter_rows(min_row=2,values_only=True):
        if row[0] is None: continue
        d[str(row[0]).zfill(4)]={"sales":float(row[6] or 0),"gp":float(row[11] or 0)}
    return d
R=load_raw("担当者別売上粗利益実績202608.xlsx")
def m(o,k): return D[o][k]["m"]["act"]
tosu={k:sum(R[c][k] for c in ["0042","0047","0046"]) for k in ["sales","gp"]}
tokyo={k:sum(R[c][k] for c in ["0550","0552"]) for k in ["sales","gp"]}
sales={"石見":m("石見","sales"),"下関":m("下関","sales"),"松江":m("松江","sales"),"鳥栖・静岡":tosu["sales"],"境港":m("境港","sales"),"東京：②":tokyo["sales"],"水産部":m("水産部","sales"),"キョウワ":round(m("キョウワ","sales"))}
sales["境港+東京"]=sales["境港"]+sales["東京：②"]; sales["営業（境港除）"]=sales["石見"]+sales["下関"]+sales["松江"]+sales["鳥栖・静岡"]
sales["合計旧タカハシ包装"]=sales["営業（境港除）"]+sales["水産部"]; sales["合計（境港含）"]=sales["合計旧タカハシ包装"]+sales["境港+東京"]
gp={"石見":m("石見","gp"),"下関":m("下関","gp"),"松江(内JF抜き）":m("松江","gp"),"鳥栖＋東京①":tosu["gp"],"東京":tokyo["gp"],"境港":m("境港","gp"),"水産部":m("水産部","gp"),"キョウワ":round(m("キョウワ","gp"))}
gp["境港＋東京"]=gp["境港"]+gp["東京"]; gp["営業（旧フクダ除）"]=gp["石見"]+gp["下関"]+gp["松江(内JF抜き）"]+gp["鳥栖＋東京①"]+gp["東京"]
gp["旧タカハシ包装合計"]=gp["営業（旧フクダ除）"]+gp["水産部"]; gp["境港含む"]=gp["旧タカハシ包装合計"]+gp["境港"]
print("SALES:",{k:round(v,3) for k,v in sales.items()}); print("GP:",{k:round(v,3) for k,v in gp.items()})
xb=xlrd.open_workbook(S+"nenkei.xls")
wb=openpyxl.Workbook(); wb.remove(wb.active)
blocks_info={}
for name,vals in [("年計表売上DATA",sales),("年計表粗利益DATA",gp)]:
    s=xb.sheet_by_name(name); ws=wb.create_sheet(name)
    for r in range(s.nrows):
        for c in range(s.ncols):
            v=s.cell_value(r,c)
            if v!="":
                if isinstance(v,float) and v==int(v) and s.cell_type(r,c)==2 and abs(v)<1e15 and r>=5 and c>=3: v=float(v)
                ws.cell(r+1,c+1).value=v
    hdr=[str(s.cell_value(5,c)) for c in range(s.ncols)]; ci={h:i for i,h in enumerate(hdr) if h}
    c8=ci["2026/8"]+1  # 1-based column of 2026/8
    ws.column_dimensions["C"].width=18
    for c in range(4,s.ncols+1): ws.column_dimensions[L(c)].width=10
    blocks=[]
    r=3
    while r<s.nrows:
        lab=s.cell_value(r,2)
        if lab and lab not in ("当月","年計","売上","粗利益") and not str(lab).startswith("20"):
            cur=r+3; nen=r+4  # 1-based rows of 当月/年計
            blocks.append((lab,cur,nen))
            if lab in vals and s.cell_value(r+2,ci["2026/7"])!="":
                ws.cell(cur,c8).value=round(vals[lab],3)
                ws.cell(nen,c8).value=f"=SUM({L(c8-11)}{cur}:{L(c8)}{cur})"
                ws.cell(cur,c8).font=Font(bold=True,color="1F4E79"); ws.cell(nen,c8).font=Font(bold=True,color="1F4E79")
        r+=1
    for row in ws.iter_rows(min_row=6):
        for cell in row:
            if isinstance(cell.value,(int,float)) or (isinstance(cell.value,str) and cell.value.startswith("=")): cell.number_format="#,##0"
    for c in range(4,s.ncols+1):
        for rr in range(1,s.nrows+1):
            cell=ws.cell(rr,c)
            if isinstance(cell.value,str) and cell.value.startswith("20") and "/" in cell.value: cell.font=Font(size=8,color="666666")
    ws.freeze_panes="D6"
    blocks_info[name]=(blocks,c8)
# legacy sheets (values only)
for name in ["売上","粗利益"]:
    s=xb.sheet_by_name(name); ws=wb.create_sheet(name)
    for r in range(s.nrows):
        for c in range(s.ncols):
            v=s.cell_value(r,c)
            if v!="": ws.cell(r+1,c+1).value=v
# charts: 年計 売上・粗利益 (直近36ヶ月) per office, on a グラフ sheet
gs=wb.create_sheet("グラフ",0)
gs["A1"]="年計推移（直近36ヶ月・単位:千円）　2026年8月更新"; gs["A1"].font=Font(bold=True,size=12)
pairs=[("石見","石見"),("下関","下関"),("松江","松江(内JF抜き）"),("鳥栖・静岡","鳥栖＋東京①"),("東京：②","東京"),("境港","境港"),("水産部","水産部"),("合計旧タカハシ包装","旧タカハシ包装合計"),("合計（境港含）","境港含む"),("キョウワ","キョウワ")]
sb,c8=blocks_info["年計表売上DATA"]; gb,_=blocks_info["年計表粗利益DATA"]
sws=wb["年計表売上DATA"]; gws=wb["年計表粗利益DATA"]
def find(blocks,lab):
    for l,cur,nen in blocks:
        if l==lab: return cur,nen
pos=0
for i,(sl,gl) in enumerate(pairs):
    scur,snen=find(sb,sl); gcur,gnen=find(gb,gl)
    c1=c8-35
    ch=LineChart(); ch.title=f"{sl}　年計（売上・粗利益）"; ch.height=7; ch.width=16
    ch.add_data(Reference(sws,min_col=c1,max_col=c8,min_row=snen,max_row=snen),from_rows=True,titles_from_data=False)
    ch.series[0].tx=None
    from openpyxl.chart.series import SeriesLabel
    ch.series[0].tx=SeriesLabel(v="年計売上")
    ch2=LineChart(); ch2.add_data(Reference(gws,min_col=c1,max_col=c8,min_row=gnen,max_row=gnen),from_rows=True,titles_from_data=False)
    ch2.series[0].tx=SeriesLabel(v="年計粗利益"); ch2.y_axis.axId=200; ch2.y_axis.crosses="max"
    ch.set_categories(Reference(sws,min_col=c1,max_col=c8,min_row=6,max_row=6))
    ch.y_axis.number_format="#,##0"; ch2.y_axis.number_format="#,##0"; ch.y_axis.title="売上"; ch2.y_axis.title="粗利益"
    ch+=ch2
    gs.add_chart(ch,f"{'A' if i%2==0 else 'K'}{3+(i//2)*15}")
wb.save(OUT+"年計表_202608.xlsx"); print("saved")
